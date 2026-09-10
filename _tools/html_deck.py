# -*- coding: utf-8 -*-
"""
No .pptx uz responsīvu HTML prezentāciju (rules_lessons.txt).

    Responsive Layouts: visiem slaidiem jāstrādā gan uz datora, gan telefonā.
    PC View:      parastais platais prezentācijas izkārtojums (nemainīgs).
    Mobile View:  @media (max-width: 768px) - elementi kārtojas viens zem otra,
                  fonti mainās plūstoši, nekas neiziet ārpus ekrāna.

Datora skats ir tas pats slaids collu precizitātē (viss pārrēķināts uz % un
cqw vienībām, tāpēc mērogojas bez JavaScript). Telefona skats to pašu saturu
pārkārto plūstošā, ritināmā lapā: kartītes viena zem otras, tabulas - pa
rindām.

Dalījumus abos skatos rāda vertikāli (skaitītājs virs saucēja).

Lietošana:
    python html_deck.py                 # visas prezentācijas /Dabaszinibas
    python html_deck.py fails.pptx      # viena prezentācija
"""

import glob
import html
import math
import re
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import mathfmt as MF
import mathfmt_prose as MP
import palette
import site_index
from courses import COURSES, DEFAULT_COURSE

EMU = 914400.0

# Katram kursam sava mape, virsraksts un failu maska ir courses.py -
# vienu un to pašu sarakstu lieto arī site_index.py.
ROOT = COURSES[DEFAULT_COURSE]["root"]

SW_IN, SH_IN = 13.333, 7.5
PT_PER_IN = 72.0


# ------------------------------------------------------------------ nolasīšana
def _rgb(color):
    # RGBColor ir tuple apakšklase - ar %-formatējumu to nedrīkst likt tieši.
    try:
        return "#" + str(color.rgb)
    except Exception:
        return None


def _fill(shape):
    try:
        f = shape.fill
        if f.type is None or int(f.type) != 1:
            return None
        return palette.slide_color(_rgb(f.fore_color), "fill")
    except Exception:
        return None


def _line(shape):
    try:
        lf = shape.line
        if lf.fill.type is None or int(lf.fill.type) != 1:
            return None, 0.0
        w = lf.width.pt if lf.width is not None else 1.0
        return palette.slide_color(_rgb(lf.color), "line"), (w or 1.0)
    except Exception:
        return None, 0.0


_ALIGN = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right",
          PP_ALIGN.JUSTIFY: "justify"}


def read_slide(slide):
    """Slaidu pārvērš vienkāršā formu sarakstā."""
    out = []
    for shp in slide.shapes:
        d = {"name": shp.name or "",
             "x": shp.left / EMU, "y": shp.top / EMU,
             "w": shp.width / EMU, "h": shp.height / EMU,
             "fill": _fill(shp), "text": None, "wrap": None,
             "round": False, "anchor": None}
        d["line"], d["lw"] = _line(shp)
        try:
            d["round"] = shp.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
        except Exception:
            pass
        if shp.has_text_frame and shp.text_frame.text.strip():
            paras = []
            for p in shp.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs)
                if not t.strip():
                    continue
                r0 = p.runs[0]
                paras.append({
                    "t": t,
                    "size": r0.font.size.pt if r0.font.size else 18.0,
                    "bold": bool(r0.font.bold),
                    "italic": bool(r0.font.italic),
                    "color": (palette.slide_color(_rgb(r0.font.color))
                              or palette.TEXT_DEFAULT),
                    "align": _ALIGN.get(p.alignment, "left"),
                    "space": p.space_before.pt if p.space_before else 0.0,
                })
            d["text"] = paras
            d["wrap"] = shp.text_frame.word_wrap
            va = shp.text_frame.vertical_anchor
            d["anchor"] = ("center" if va == MSO_ANCHOR.MIDDLE else
                           "flex-end" if va == MSO_ANCHOR.BOTTOM else
                           "flex-start")
        out.append(d)
    return _lighten_on_dark(out)


def _lighten_on_dark(shapes):
    """Tekstam tumšā laukumā iedod gaišo krāsu variantu.

    Ģeometrija ir šejienes ziņa, krāsu izvēle - palette.on_dark ziņa.
    """
    dark = [s for s in shapes
            if s["fill"] and s["text"] is None and is_dark(s["fill"])
            and "MATH:" not in s["name"] and "|ARR:" not in s["name"]
            and "|SEG:" not in s["name"]]
    for t in shapes:
        if not t["text"]:
            continue
        for b in dark:
            if (b["x"] <= t["x"] + .05 and b["y"] <= t["y"] + .05
                    and b["x"] + b["w"] >= t["x"] + t["w"] - .05
                    and b["y"] + b["h"] >= t["y"] + t["h"] - .05):
                for p in t["text"]:
                    p["color"] = palette.on_dark(p["color"])
                break
    return shapes


# --------------------------------------------------------------------- HTML
def esc(s):
    """Tikai HTML rakstzīmju aizsegšana - bez formulu noformējuma."""
    return html.escape(s, quote=False)


def txt_html(s):
    """Teksts HTML: aizsegts un ar uzzīmētām vektoru bultiņām.

    Viss slaida teksts iet caur šo funkciju (DRY) - tāpēc vektora
    pieraksts izskatās vienādi virsrakstos, kartītēs, tabulās un formulās.
    """
    if not MF.has_vector(s):
        return esc(s)
    return "".join(esc(v) if k == "t" else vec_html(v)
                   for k, v in MF.split_vectors(s))


def vec_html(base):
    """Vektora simbols: bultiņu zīmē CSS, nevis fonta kombinējošā zīme."""
    return '<span class="vv">%s</span>' % esc(base)


def _sp(s):
    """Atstarpes HTML nesaspiež - tās notur formulu atstatumus.

    Rindas sākuma un beigu atstarpi pārlūks izmet pavisam, tāpēc centrēts
    gabals ("√2500" un " = 50 N") saslīdētu kopā; te tās paliek kā &nbsp;.
    """
    lead = len(s) - len(s.lstrip(" "))
    trail = len(s) - len(s.rstrip(" ")) if s.strip() else 0
    core = s[lead:len(s) - trail] if trail else s[lead:]
    body = txt_html(core).replace("  ", "&nbsp;&nbsp;")
    return "&nbsp;" * lead + body + "&nbsp;" * trail


# Saknes zīmes augstums fonta izmēra daļās nāk no mathfmt - tie paši mēri,
# pēc kuriem zīmi uzzīmē arī .pptx, tāpēc abi skati sakrīt.
ROOT_EM = MF.root_em([("t", "")])
ROOT_EM_TALL = MF.root_em([("f", "", "")])


def root_svg(h=ROOT_EM, cls="rk", style=""):
    """Saknes zīme kā SVG - viena forma abiem skatiem (MF.root_pts)."""
    w, _ = MF.root_pts(h)
    return ('<svg class="%s"%s viewBox="0 0 %.4f %.4f" '
            'preserveAspectRatio="none" aria-hidden="true">'
            '<path d="%s"/></svg>'
            % (cls, (' style="%s"' % style) if style else "",
               w, h, MF.root_path_d(h)))


def root_html(inner, tall=False):
    """√-izteiksme: vinkuls (svītra) pāri VISAI izteiksmei, ne tikai iekavai.

    Saknes zīmi zīmē pati lapa (MF.ROOT_PTS), tāpēc tā izstiepjas līdz
    satura augstumam - arī tad, ja zem vinkula ir vertikāla daļa - un
    vinkuls turpinās no tās augšmalas. `inner` jau ir gatavs HTML.
    """
    h = ROOT_EM_TALL if tall else ROOT_EM
    w, _ = MF.root_pts(h)
    return ('<span class="rt" style="--rw:%.3fem">%s'
            '<span class="rv">%s</span></span>'
            % (w, root_svg(h), inner))


def frac_span(num, den):
    """Vertikāla daļa: skaitītājs virs saucēja."""
    return ('<span class="f"><span class="n">%s</span>'
            '<span class="d">%s</span></span>'
            % (txt_html(num), txt_html(den)))


def atoms_html(atoms):
    """Atomu virkne uz HTML - viena vieta abiem parsētājiem (DRY)."""
    out = []
    for a in atoms:
        if a[0] == "t":
            out.append(_sp(a[1]))
        elif a[0] == "r":
            out.append(root_html(atoms_html(a[1]),
                                 tall=any(x[0] == "f" for x in a[1])))
        else:
            out.append(frac_span(a[1], a[2]))
    return "".join(out)


def frac_html(text):
    """Parasta teksta dalījumi HTML - vertikāla daļa kā <span class=f>.

    Lieto piesardzīgo atpazīšanu: mērvienības (m/s, kg/m³) un vārdu pāri
    (garums/augstums) paliek rindā, bet 1/16, F/S, 1/r² kļūst vertikāli.
    """
    return atoms_html(MP.parse_prose(text))


# ---------------------------------------------------------------- zīmējumi
# Bultas un palīglīnijas .pptx failā ir formas, kuru nosaukumā ierakstīta
# ģeometrija ("FIG7|ARR:x1,y1,x2,y2,lw,head"). Tāpēc HTML var uzzīmēt tieši
# to pašu, nevis minēt pēc formas rāmja.

def fig_part(s):
    """Formas nosaukumu pārvērš zīmējuma daļā vai atgriež None."""
    nm = s["name"]
    if "|" not in nm:
        return None
    fid, tail = nm.split("|", 1)
    if tail.startswith(("ARR:", "SEG:")):
        try:
            x1, y1, x2, y2, lw, head = (float(v)
                                        for v in tail[4:].split(","))
        except ValueError:
            return None
        return {"fid": fid, "kind": "line", "x1": x1, "y1": y1,
                "x2": x2, "y2": y2, "lw": lw, "head": bool(head),
                "color": s["fill"] or s["line"] or palette.TEXT_DEFAULT}
    if tail == "LBL":
        return {"fid": fid, "kind": "label", "shape": s}
    return None


def _arrow_path(p, k=1.0):
    """Bultas kontūra SVG ceļš; k - mērogs no collām uz zīmējuma vienībām."""
    dx, dy = p["x2"] - p["x1"], p["y2"] - p["y1"]
    ln = math.hypot(dx, dy) or 1e-6
    ux, uy = dx / ln, dy / ln
    px, py = -uy, ux
    hl = min(0.20, ln * 0.34) if p["head"] else 0.0
    hw, sw = p["lw"] * 2.2, p["lw"] / 2.0

    def pt(u, v):
        return ((p["x1"] + ux * u + px * v) * k,
                (p["y1"] + uy * u + py * v) * k)

    if not p["head"]:
        pts = [pt(0, -sw), pt(ln, -sw), pt(ln, sw), pt(0, sw)]
    else:
        pts = [pt(0, -sw), pt(ln - hl, -sw), pt(ln - hl, -hw), pt(ln, 0),
               pt(ln - hl, hw), pt(ln - hl, sw), pt(0, sw)]
    return "M%s Z" % " L".join("%.3f %.3f" % q for q in pts)


def pt2cqw(pt):
    """Punktus pārvērš cqw vienībās - slaida platums ir 100 cqw."""
    return pt / PT_PER_IN / SW_IN * 100.0


def _style_box(s, fill=True):
    st = ["left:%.4f%%" % (s["x"] / SW_IN * 100),
          "top:%.4f%%" % (s["y"] / SH_IN * 100),
          "width:%.4f%%" % (s["w"] / SW_IN * 100),
          "height:%.4f%%" % (s["h"] / SH_IN * 100)]
    if fill and s["fill"]:
        st.append("background:%s" % s["fill"])
    if s["line"]:
        st.append("border:%.3fcqw solid %s" % (pt2cqw(s["lw"]), s["line"]))
    return ";".join(st)


def stage_svg(shapes):
    """Visas slaida bultas vienā SVG pārklājumā (slaida koordinātēs)."""
    parts = [p for p in (fig_part(s) for s in shapes)
             if p and p["kind"] == "line"]
    if not parts:
        return ""
    body = "".join('<path d="%s" fill="%s"/>'
                   % (_arrow_path(p, 100.0), p["color"]) for p in parts)
    return ('<svg class="fg" viewBox="0 0 %.1f %.1f" '
            'preserveAspectRatio="none">%s</svg>'
            % (SW_IN * 100, SH_IN * 100, body))


def stage_html(shapes):
    """Datora skats: precīzs slaida atveids (mērogojas pats, bez JS)."""
    out = []
    for s in shapes:
        extra = ""
        if s["round"] and min(s["w"], s["h"]) > 0.08:
            extra = ";border-radius:%.3fcqw" % (min(s["w"], s["h"]) * 0.05
                                                / SW_IN * 100)
        if s["text"] is None:
            if "|ARR:" in s["name"] or "|SEG:" in s["name"]:
                continue                         # bultas zīmē stage_svg()
            role = s["name"].rsplit(":", 1)[-1]
            if role.startswith("rg"):            # saknes zīme
                out.append(root_svg(
                    ROOT_EM_TALL if role.endswith("T") else ROOT_EM, "rg",
                    "%s;fill:%s" % (_style_box(s, fill=False), s["fill"])))
            else:
                out.append('<i style="%s%s"></i>' % (_style_box(s), extra))
            continue
        css = ";justify-content:%s" % (s["anchor"] or "flex-start")
        if s["wrap"] is False:
            css += ";white-space:nowrap;align-items:center"
        out.append('<div class="tb" style="%s%s%s">'
                   % (_style_box(s), extra, css))
        for p in s["text"]:
            out.append(
                '<p style="font-size:%.4fcqw;color:%s;text-align:%s;'
                'padding-top:%.4fcqw%s%s">%s</p>'
                % (pt2cqw(p["size"]), p["color"], p["align"],
                   pt2cqw(p["space"]),
                   ";font-weight:700" if p["bold"] else "",
                   ";font-style:italic" if p["italic"] else "",
                   _sp(p["t"])))
        out.append("</div>")
    out.append(stage_svg(shapes))
    return "\n".join(out)


# ------------------------------------------------- telefona skats (pārkārtots)
def _group(name):
    """'BOX:PANEL7' -> ('BOX', 'PANEL7')"""
    if not name:
        return "", ""
    n = name.split("|")[0]
    if ":" in n:
        a, b = n.split(":", 1)
        return a, b
    return n, ""


GAP = ""


def _math_runs(shapes):
    """Formulu gabalus (MATH:...) saliek atpakaļ vienā rindā ar daļām."""
    runs, order = {}, []
    for s in shapes:
        nm = s["name"]
        if "MATH:" not in nm:
            continue
        owner = nm.split("|")[0] if "|" in nm else ""
        tail = nm.split("MATH:", 1)[1]
        mid, role = tail.split(":", 1)
        if mid not in runs:
            runs[mid] = {"parts": [], "owner": owner}
            order.append(mid)
        txt = s["text"][0]["t"] if s["text"] else ""
        runs[mid]["parts"].append((role, txt, s))

    built = {}
    for mid in order:
        r = runs[mid]
        # Lomu priedēklis "r" nozīmē "zem saknes vinkula": rt, rn, rd, rb.
        # Zīme (rg) atver sakni, vinkuls (rv) to aizver - tā saknes saturs
        # telefona skatā saliekas atpakaļ tieši tāds pats kā slaidā.
        pieces, num, root, root_tall = [], None, None, False
        for role, txt, s in r["parts"]:
            here = pieces if root is None else root
            if role.startswith("rg"):
                root, root_tall = [], role.endswith("T")
            elif role == "rv":
                pieces.append(root_html("".join(root or []), root_tall))
                root = None
            elif role in ("b", "rb"):
                continue                      # daļas svītru uzzīmē CSS
            elif role in ("t", "rt"):
                # 2+ atstarpes atdala patstāvīgas formulas vienā rindā -
                # telefonā tās drīkst pārlēkt uz nākamo rindu veselas
                here.append(GAP.join(txt_html(x) for x in
                                     re.split(r"\s{2,}", txt)))
            elif role in ("n", "rn"):
                num = txt
            elif role in ("d", "rd"):
                here.append(frac_span(num or "", txt))
        first = next((s for _, _, s in r["parts"] if s["text"]), None)
        p0 = first["text"][0] if first else {}
        chunks = [c for c in "".join(pieces).split(GAP) if c.strip()]
        built[mid] = {
            "html": "".join('<span class="mc">%s</span>' % c for c in chunks),
            "owner": r["owner"],
            "y": min(p[2]["y"] for p in r["parts"]),
            "x": min(p[2]["x"] for p in r["parts"]),
            "size": p0.get("size", 18.0),
            "color": p0.get("color", palette.TEXT_DEFAULT),
            "bold": p0.get("bold", False),
        }
    return built


def _ptcls(size):
    if size >= 30:
        return "h1"
    if size >= 22:
        return "h2"
    if size >= 18:
        return "p1"
    if size >= 14:
        return "p2"
    return "p3"


def _pstyle(p):
    st = []
    if p["color"] != palette.TEXT_DEFAULT:
        st.append("color:%s" % p["color"])
    if p["bold"]:
        st.append("font-weight:700")
    if p["italic"]:
        st.append("font-style:italic")
    return ' style="%s"' % ";".join(st) if st else ""


def flow_bg(shapes):
    """Pilnas lapas fons (titullapa, uzdevumu šķirlapa) - lai telefona
    skatā gaišais teksts nepaliktu uz balta."""
    bg = None
    for s in shapes:
        if (s["text"] is None and s["fill"]
                and s["w"] > SW_IN * 0.95 and s["h"] > SH_IN * 0.80):
            bg = s["fill"]
    return bg


def is_dark(hexcolor):
    try:
        r, g, b = (int(hexcolor[i:i + 2], 16) for i in (1, 3, 5))
    except Exception:
        return False
    return (0.299 * r + 0.587 * g + 0.114 * b) < 140


def collect_figs(shapes):
    """Zīmējuma daļas sagrupē pa zīmējumiem (fid)."""
    figs = {}
    for s in shapes:
        p = fig_part(s)
        if not p:
            continue
        f = figs.setdefault(p["fid"], {"lines": [], "labels": [], "y": s["y"],
                                       "x": s["x"]})
        f["y"] = min(f["y"], s["y"])
        f["x"] = min(f["x"], s["x"])
        f["lines" if p["kind"] == "line" else "labels"].append(
            p if p["kind"] == "line" else s)
    return figs


def fig_html(f):
    """Zīmējums telefona skatā: tas pats attēls, tikai lapas platumā."""
    xs, ys = [], []
    for p in f["lines"]:
        for a, b, r in ((p["x1"], p["y1"], p["lw"] * 2.4),
                        (p["x2"], p["y2"], p["lw"] * 2.4)):
            xs += [a - r, a + r]
            ys += [b - r, b + r]
    for s in f["labels"]:
        xs += [s["x"], s["x"] + s["w"]]
        ys += [s["y"], s["y"] + s["h"]]
    if not xs:
        return ""
    pad = 0.06
    x0, x1 = min(xs) - pad, max(xs) + pad
    y0, y1 = min(ys) - pad, max(ys) + pad
    w, h = max(x1 - x0, 0.01), max(y1 - y0, 0.01)

    body = "".join('<path d="%s" fill="%s"/>'
                   % (_arrow_path(dict(p, x1=p["x1"] - x0, y1=p["y1"] - y0,
                                       x2=p["x2"] - x0, y2=p["y2"] - y0),
                                  100.0), p["color"])
                   for p in f["lines"])
    out = ['<div class="fig" style="aspect-ratio:%.3f/%.3f">' % (w, h),
           '<svg viewBox="0 0 %.1f %.1f" preserveAspectRatio="none">%s</svg>'
           % (w * 100, h * 100, body)]
    for s in f["labels"]:
        pr = s["text"][0]
        out.append('<b style="left:%.3f%%;top:%.3f%%;width:%.3f%%;'
                   'font-size:%.3fcqw;color:%s">%s</b>'
                   % ((s["x"] - x0) / w * 100, (s["y"] - y0) / h * 100,
                      s["w"] / w * 100, pr["size"] / PT_PER_IN / w * 100,
                      pr["color"], txt_html(pr["t"])))
    out.append("</div>")
    return "".join(out)


def flow_html(shapes):
    """Telefona skats: tas pats saturs vienā plūstošā kolonnā."""
    math = _math_runs(shapes)
    figs = collect_figs(shapes)

    boxes = {}                       # 'BOX:PANEL7' -> rāmja forma
    for s in shapes:
        if s["text"] is None and s["name"].startswith("BOX:"):
            boxes.setdefault(s["name"], s)

    # --- tabulas ----------------------------------------------------------
    tables = {}
    for s in shapes:
        kind, rest = _group(s["name"])
        if kind not in ("TBLHT", "TBLCT", "TBLH", "TBLC"):
            continue
        tid = rest.split(":")[0]
        t = tables.setdefault(tid, {"hdr": {}, "rows": {}, "y": s["y"]})
        t["y"] = min(t["y"], s["y"])
        if s["text"] is None:
            continue
        if kind == "TBLHT":
            t["hdr"][int(rest.split(":")[1])] = txt_html(
                s["text"][0]["t"])
        elif kind == "TBLCT":
            _, i, j = rest.split(":")
            t["rows"].setdefault(int(i), {})[int(j)] =                 frac_html(s["text"][0]["t"])

    for m in math.values():                     # formulas tabulu šūnās
        own = m["owner"]
        if own.startswith("TBLCT:"):
            tid, i, j = own.split(":", 1)[1].split(":")
            if tid in tables:
                tables[tid]["rows"].setdefault(int(i), {})[int(j)] = m["html"]

    # --- vienumi plūsmā ---------------------------------------------------
    items, done_tbl = [], set()
    for s in shapes:
        kind, gid = _group(s["name"])
        if kind in ("TBLH", "TBLC", "TBLHT", "TBLCT"):
            tid = gid.split(":")[0]
            if tid not in done_tbl:
                done_tbl.add(tid)
                items.append((tables[tid]["y"], s["x"], None, "table", tid))
            continue
        if "MATH:" in s["name"] or s["name"] == "RULE":
            continue
        if fig_part(s):                      # zīmējuma daļas - kopā, zemāk
            continue
        if s["text"] is None or s["name"] == "FOOTER":
            continue
        g = gid if kind == "TXT" else None
        items.append((s["y"], s["x"], g, "text", s))

    for m in math.values():                     # formulas ārpus tabulām
        if m["owner"].startswith("TBLCT:"):
            continue
        k, g = _group(m["owner"])
        items.append((m["y"], m["x"], g if k == "TXT" else None, "math", m))

    for f in figs.values():
        items.append((f["y"], f["x"], None, "fig", f))

    # Uzdevumu slaidos telefona skatā jāsaglabā risinājuma secība
    # (Dots -> Jāaprēķina -> Formulas -> Aprēķins -> Atbilde), nevis
    # jāseko slaida kolonnām.
    ORD = {"SOLVE0": 2.30, "SOLVE1": 2.31, "SOLVE2": 2.32,
           "CALC": 2.40, "CARD": 2.50}

    def _key(it):
        y, x, gid = it[0], it[1], it[2]
        return (round(ORD.get(gid, y), 2), 0.0 if gid in ORD else x)

    items.sort(key=_key)

    SENT = object()
    out, cur = [], SENT
    for y, x, gid, kind, obj in items:
        if gid != cur or gid is None:
            if cur is not SENT:
                out.append("</div>")
            cur = gid
            card = boxes.get("BOX:" + gid) if gid else None
            if card is not None:
                st = []
                if card["fill"] and card["fill"] != "#FFFFFF":
                    st.append("background:%s" % card["fill"])
                if card["line"]:
                    st.append("border-color:%s" % card["line"])
                out.append('<div class="card"%s>'
                           % (' style="%s"' % ";".join(st) if st else ""))
            else:
                cls = "blk"
                if kind == "text" and obj["name"] in ("KICKER", "TITLE",
                                                      "BADGE"):
                    cls = obj["name"].lower()
                out.append('<div class="%s">' % cls)

        if kind == "table":
            t = tables[obj]
            out.append('<div class="mtbl">')
            for i in sorted(t["rows"]):
                row = t["rows"][i]
                out.append('<div class="mrow">')
                for j in sorted(row):
                    out.append('<div class="mcell">'
                               '<span class="ml">%s</span>'
                               '<span class="mv">%s</span></div>'
                               % (t["hdr"].get(j, ""), row[j]))
                out.append("</div>")
            out.append("</div>")
        elif kind == "fig":
            out.append(fig_html(obj))
        elif kind == "math":
            out.append('<p class="mf" style="color:%s">%s</p>'
                       % (obj["color"], obj["html"]))
        else:
            for pr in obj["text"]:
                out.append('<p class="%s"%s>%s</p>'
                           % (_ptcls(pr["size"]), _pstyle(pr),
                              frac_html(pr["t"])))
    if cur is not SENT:
        out.append("</div>")
    return "\n".join(out)


# ---------------------------------------------------------------------- lapa
CSS = """
*{box-sizing:border-box}
html,body{margin:0;padding:0}
body{background:var(--bg);color:var(--fg);
     font-family:var(--font);
     -webkit-text-size-adjust:100%}
.top{position:sticky;top:0;z-index:20;display:flex;gap:.6rem;
     align-items:center;flex-wrap:wrap;
     padding:.6rem .9rem;background:var(--grad);color:#fff;
     box-shadow:var(--sh-md)}
.top b{font-size:1rem;font-weight:600;font-family:var(--font-h)}
.top .sub{font-size:.78rem;color:rgba(255,255,255,.82);flex:1 1 12rem}
.top button{font:inherit;font-size:.85rem;font-weight:500;border:0;
     border-radius:var(--r-pill);padding:.4rem .95rem;color:#fff;
     background:rgba(255,255,255,.18);cursor:pointer;
     transition:background .15s}
.top button:hover{background:rgba(255,255,255,.30)}
.top .cnt{font-size:.8rem;color:rgba(255,255,255,.82);
     font-variant-numeric:tabular-nums}
.top .back{font-size:.82rem;color:#fff;text-decoration:none;
     padding:.35rem .75rem;border-radius:var(--r-pill);
     background:rgba(255,255,255,.14);transition:background .15s}
.top .back:hover{background:rgba(255,255,255,.28)}

.deck{padding:.9rem;display:flex;flex-direction:column;gap:.9rem}
.slide{container-type:inline-size;background:var(--surface);
       border-radius:var(--r);overflow:hidden;box-shadow:var(--sh-md);
       position:relative}
.num{position:absolute;z-index:5;right:.6rem;top:.4rem;font-size:.7rem;
     color:var(--dim);font-variant-numeric:tabular-nums}

/* ---------- DATORA SKATS: precīzs slaids ---------- */
.stage{position:relative;width:100%;aspect-ratio:13.333/7.5;
       background:var(--surface)}
.stage i,.stage .rg{position:absolute;display:block}
/* vektoru zīmējums: bultas vienā SVG pāri slaidam */
.stage .fg{position:absolute;left:0;top:0;width:100%;height:100%;
           overflow:visible;pointer-events:none}
.fig{position:relative;container-type:inline-size;width:100%;
     margin:.6rem 0}
.fig svg{position:absolute;left:0;top:0;width:100%;height:100%}
.fig b{position:absolute;text-align:center;font-weight:600;
       line-height:1.1;transform:translateY(-.1em)}
.stage .tb{position:absolute;display:flex;flex-direction:column;
           justify-content:flex-start;overflow:visible}
.stage .tb p{margin:0;line-height:1.22}
.flow{display:none}

/* ---------- vertikāla daļa (a/b) ---------- */
.f{display:inline-flex;flex-direction:column;align-items:stretch;
   text-align:center;vertical-align:middle;margin:0 .28em;line-height:1.14}
.f .n{display:block;padding:0 .2em}
.f .d{display:block;padding:0 .2em;border-top:.075em solid currentColor}

/* ---------- kvadrātsakne: vinkuls pāri visai izteiksmei ----------
   Saknes zīme ir zīmēta (SVG), tāpēc tā izstiepjas līdz izteiksmes
   augstumam un vinkuls turpinās tieši no tās augšējā stūra. */
.rt{position:relative;display:inline-block;white-space:nowrap;
    --rw:.66em;padding-left:var(--rw);margin:0 .06em}
.rt .rk{position:absolute;left:0;top:0;bottom:0;width:var(--rw);height:auto;
        fill:currentColor}
.rt .rv{display:inline-block;border-top:.075em solid currentColor;
        padding:.16em .2em 0 .06em}

/* ---------- vektors: bultiņa virs simbola ----------
   Unikoda kombinējošā bultiņa (U+20D7) lapas fontos vai nu iztrūkst, vai
   nostājas blakus burtam, tāpēc bultiņu zīmē CSS: kāts pāri simbolam un
   trīsstūra uzgalis labajā galā. */
.vv{position:relative;display:inline-block;line-height:1}
.vv::before{content:"";position:absolute;left:-.03em;right:.13em;top:-.05em;
            border-top:.07em solid currentColor}
.vv::after{content:"";position:absolute;right:-.04em;top:-.105em;
           width:0;height:0;border-left:.2em solid currentColor;
           border-top:.1em solid transparent;
           border-bottom:.1em solid transparent}

/* ---------- TELEFONA SKATS ---------- */
@media (max-width:768px){
  .deck{padding:.55rem;gap:.7rem}
  .stage{display:none}
  .flow{display:block;padding:.85rem .8rem 1rem}
  .slide{border-radius:.45rem}
  .num{position:static;display:block;text-align:right;padding:.4rem .8rem 0}

  .flow p{margin:.28rem 0;line-height:1.36;overflow-wrap:anywhere}
  .flow .h1{font-size:clamp(1.15rem,5.6vw,1.7rem);font-weight:700;
            line-height:1.22}
  .flow .h2{font-size:clamp(1.02rem,4.7vw,1.4rem);font-weight:700;
            line-height:1.26}
  .flow .p1{font-size:clamp(.95rem,4.1vw,1.15rem)}
  .flow .p2{font-size:clamp(.88rem,3.7vw,1.05rem)}
  .flow .p3{font-size:clamp(.8rem,3.3vw,.95rem)}
  .flow .kicker p{font-size:clamp(.68rem,3vw,.8rem);letter-spacing:.06em;
                  text-transform:uppercase;color:var(--dim)}
  .flow .title p{font-size:clamp(1.15rem,5.4vw,1.6rem);font-weight:600;
                 font-family:var(--font-h);color:var(--primary)}
  .flow .badge p{font-size:clamp(.78rem,3.4vw,.95rem);font-weight:600;
                 color:var(--amber-ink)}
  .flow.dark p{color:#EEF2FF}
  .flow.dark .card{background:rgba(255,255,255,.07);
                    border-color:rgba(255,255,255,.25)}
  .flow .card{border:1px solid var(--line);border-left-width:4px;
              border-left-color:var(--violet);border-radius:.75rem;
              padding:.65rem .75rem;margin:.55rem 0;
              background:var(--surface)}
  .flow .blk{margin:.35rem 0}
  .flow .mf{font-size:clamp(1.15rem,6vw,1.75rem);font-weight:600;
            text-align:center;color:var(--primary);margin:.5rem 0}
  .flow .mf .mc{display:inline-block;margin:.15rem .5rem}
  .flow .mfcard{background:var(--surface2);border-radius:.75rem}

  .flow .mtbl{margin:.55rem 0;display:flex;flex-direction:column;gap:.5rem}
  .flow .mrow{border:1px solid var(--line);border-radius:.75rem;
              overflow:hidden}
  .flow .mcell{display:flex;gap:.6rem;align-items:baseline;
               padding:.45rem .65rem;border-top:1px solid var(--line)}
  .flow .mrow .mcell:first-child{border-top:0;background:var(--surface2);
                                 font-weight:600;color:var(--primary)}
  .flow .ml{flex:0 0 38%;font-size:clamp(.7rem,3vw,.82rem);
            color:var(--dim);text-transform:uppercase;
            letter-spacing:.06em}
  .flow .mv{flex:1 1 auto;font-size:clamp(.88rem,3.8vw,1.05rem)}
}

/* ---------- PILNEKRĀNA SKATS: viens slaids uz visa ekrāna ---------- */
/* Slaids patur 13.333:7.5 malu attiecību un ieņem tik daudz, cik ietilpst;
   pārējais paliek tumšs. Teksts mērogojas pats, jo .slide ir konteiners. */
body.fs{background:#1E1B4B;overflow:hidden}
body.fs .top{display:none}
body.fs .deck{position:fixed;inset:0;margin:0;padding:0;gap:0;
              align-items:center;justify-content:center}
body.fs .slide{display:none;margin:0;border-radius:0;box-shadow:none;
               width:min(100vw,calc(100vh * 13.333 / 7.5));
               height:min(100vh,calc(100vw * 7.5 / 13.333))}
body.fs .slide.on{display:block}
body.fs .num{display:none}
/* pilnekrānā vienmēr rāda īsto slaidu, arī telefonā */
body.fs .stage{display:block}
body.fs .flow{display:none}

.fsui{display:none}
body.fs .fsui{display:flex;position:fixed;z-index:30;left:50%;
  bottom:calc(1rem + env(safe-area-inset-bottom));transform:translateX(-50%);
  align-items:center;gap:.35rem;padding:.3rem .45rem;color:#fff;
  border-radius:var(--r-pill);background:rgba(15,12,45,.7);
  opacity:.45;transition:opacity .2s}
body.fs .fsui:hover,body.fs .fsui:focus-within{opacity:1}
.fsui button{font:inherit;line-height:1;border:0;cursor:pointer;color:#fff;
  background:transparent;border-radius:var(--r-pill);padding:.3rem .7rem;
  transition:background .15s}
.fsui button:hover{background:rgba(255,255,255,.18)}
.fsui .nav{font-size:1.25rem;padding:.15rem .6rem}
.fsui .cnt{font-size:.8rem;min-width:4.2rem;text-align:center;
           font-variant-numeric:tabular-nums}

/* uz šauriem ekrāniem neko nedrīkst griezt horizontāli */
@media (max-width:768px){
  html,body{overflow-x:hidden}
  .top button{font-size:.78rem;padding:.35rem .7rem}
  .top b{font-size:.95rem;flex:1 1 100%;order:2}
  .top .back{order:1}
  .flow *{max-width:100%}
}
"""

JS = """
(function(){
  /* Katrai daļai viens uzdevums: ekrāns (fullscreen API), stāvoklis
     (kurš slaids) un attēlošana. Taustiņi un pogas tikai sauc tos pašus
     trīs darbus - open, close, go. */
  var slides=[].slice.call(document.querySelectorAll('.slide'));
  var cnt=document.getElementById('cnt');
  var fscnt=document.getElementById('fscnt');
  var deck=document.getElementById('deck');
  var cur=0,on=false;

  /* --- ekrāns --- */
  var root=document.documentElement;
  function fsel(){return document.fullscreenElement||
                         document.webkitFullscreenElement;}
  function ask(){
    var f=root.requestFullscreen||root.webkitRequestFullscreen;
    if(f){try{var p=f.call(root);if(p&&p.catch)p.catch(function(){});}
          catch(e){}}
  }
  function drop(){
    var f=document.exitFullscreen||document.webkitExitFullscreen;
    if(fsel()&&f){try{var p=f.call(document);if(p&&p.catch)p.catch(function(){});}
                  catch(e){}}
  }

  /* --- attēlošana --- */
  function render(){
    for(var i=0;i<slides.length;i++)
      slides[i].classList.toggle('on',i===cur);
    var pos=(cur+1)+' / '+slides.length;
    fscnt.textContent=pos;
    cnt.textContent=on?pos:slides.length+' slaidi';
  }

  /* --- stāvoklis --- */
  function go(i){
    cur=i<0?0:(i>=slides.length?slides.length-1:i);
    render();
  }
  function open(){
    on=true;document.body.classList.add('fs');ask();render();
  }
  function close(){
    on=false;document.body.classList.remove('fs');drop();render();
    if(slides[cur])slides[cur].scrollIntoView({block:'start'});
  }

  /* --- vadība --- */
  document.getElementById('fs').addEventListener('click',open);
  document.getElementById('fsx').addEventListener('click',close);
  document.getElementById('prev').addEventListener('click',function(e){
    e.stopPropagation();go(cur-1);});
  document.getElementById('next').addEventListener('click',function(e){
    e.stopPropagation();go(cur+1);});
  deck.addEventListener('click',function(){if(on)go(cur+1);});

  document.addEventListener('keydown',function(e){
    if(!on)return;
    var k=e.key;
    if(k==='ArrowRight'||k===' '||k==='PageDown'||k==='Enter'){go(cur+1);}
    else if(k==='ArrowLeft'||k==='PageUp'||k==='Backspace'){go(cur-1);}
    else if(k==='Home'){go(0);}
    else if(k==='End'){go(slides.length-1);}
    else if(k==='Escape'){close();}
    else{return;}
    e.preventDefault();
  });
  /* Esc pilnekrānā pārtver pārlūks - tad jāsakārto pašiem. */
  function synced(){if(on&&!fsel())close();}
  document.addEventListener('fullscreenchange',synced);
  document.addEventListener('webkitfullscreenchange',synced);

  render();
})();
"""

PAGE = """<!DOCTYPE html>
<html lang="lv">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<title>%(title)s</title>
%(fonts)s<style>%(root)s%(css)s</style>
</head>
<body>
<div class="top">
  <a class="back" href="../index.html">&#8592; Uz t&#275;m&#257;m</a>
  <b>%(title)s</b>
  <span class="sub">%(kicker)s</span>
  <span class="cnt" id="cnt"></span>
  <button id="fs" type="button">Pilnekr&#257;ns</button>
</div>
<div class="deck" id="deck">
%(slides)s
</div>
<div class="fsui" id="fsui">
  <button class="nav" id="prev" type="button" aria-label="Iepriek&#353;&#275;jais">&#8249;</button>
  <span class="cnt" id="fscnt"></span>
  <button class="nav" id="next" type="button" aria-label="N&#257;kamais">&#8250;</button>
  <button id="fsx" type="button">Esc</button>
</div>
<script>%(js)s</script>
</body>
</html>
"""


def convert(pptx_path, out_path=None):
    prs = Presentation(pptx_path)
    title = os.path.splitext(os.path.basename(pptx_path))[0]
    kicker = ""
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        shapes = read_slide(slide)
        if not kicker:
            for s in shapes:
                if s["name"] == "KICKER" and s["text"]:
                    kicker = s["text"][0]["t"]
                    break
        bg = flow_bg(shapes)
        cls = "flow dark" if bg and is_dark(bg) else "flow"
        st = ' style="background:%s"' % bg if bg else ""
        parts.append(
            '<section class="slide">\n<span class="num">%d</span>\n'
            '<div class="stage">\n%s\n</div>\n'
            '<div class="%s"%s>\n%s\n</div>\n</section>'
            % (i, stage_html(shapes), cls, st, flow_html(shapes)))

    page = PAGE % {"title": esc(title), "kicker": esc(kicker),
                   "css": CSS, "js": JS, "slides": "\n".join(parts),
                   "root": palette.root_css(), "fonts": palette.FONT_LINK}
    if out_path is None:
        out_path = os.path.splitext(pptx_path)[0] + ".html"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(page)
    return out_path, len(parts)


def build_index(root=ROOT, title=None):
    """Novecojis vārds - sarakstus tagad būvē site_index.py."""
    return site_index.build_course_index(site_index.slug_for_root(root))


def main(argv):
    """html_deck.py [kurss] [faili...]

    Bez argumentiem - viss noklusētais kurss; "fizika" vai "dabaszinibas" -
    viss attiecīgais kurss; faili - tikai tie (indeksu tad nepārbūvē).
    """
    course = DEFAULT_COURSE
    args = argv[1:]
    if args and args[0] in COURSES:
        course = args.pop(0)
    c = COURSES[course]
    targets = args or sorted(glob.glob(os.path.join(c["root"], "*",
                                                    c["pattern"])))
    n = 0
    for p in targets:
        out, k = convert(p)
        n += 1
        print("%3d slaidi  %s" % (k, os.path.basename(out)))
    print("Kopā: %d prezentācijas" % n)
    if not args:
        print("Saraksts:   %s" % site_index.build_course_index(course))
        print("Sākumlapa:  %s" % site_index.build_home())


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv)
