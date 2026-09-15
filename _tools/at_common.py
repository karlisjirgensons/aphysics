# -*- coding: utf-8 -*-
"""
Ātro testu (ĀT) veidotājs — viena lapa, viens variants, atbildes A/B/C/D.

Atšķirībā no pārbaudes darba (pd_common.py) šeit nav ne aprēķinu, ne brīvo
atbilžu: skolēns katram jautājumam apvelk vienu burtu atbilžu tabulā lapas
apakšā. Darbs paredzēts ~15 minūtēm stundas sākumā vai beigās.

No viena satura apraksta (skat. at_*.py) uzbūvē divus .docx failus:
  «<nosaukums>.docx»             — skolēna lapa (viena A4 lapa)
  «<nosaukums> (atbildes).docx»  — atbilžu atslēga un vērtēšanas skala

Noformējums ir tas pats, kas PD lapām: Calibri, tumši zila galvene, ATGĀDNE
kaste, pelēki kursīva paskaidrojumi.

Satura shēma
------------
AT = {
    "nr": 1, "nosaukums": ..., "mape": ..., "fails": ...,
    "prieksmets": None | "Fizika I  |  10. klase",
    "sakne": None | "Fizika_1",
    "stundas": "1.5.-1.10.", "datums": "07.10.2026.", "laiks": 15,
    "apraksts": ..., "atgadne": [...],
    "struktura": [(jautājumi, sasniedzamais rezultāts, stunda), ...],
    "jautajumi": [(teksts, [4 atbildes], pareizās atbildes indekss), ...],
}
"""

import os

from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

import pd_common as C

AILES = 6                    # cik jautājumu vienā atbilžu tabulas blokā
BURTU_ATSTARPE = "     "     # starp A B C D atbilžu tabulā
JAUT_PT = 10.0               # jautājumu un atbilžu fonta izmērs
ATSTARPE = "      "          # starp atbilžu variantiem vienā rindā


# -------------------------------------------------------------------- palīgi
def _konteksts(at):
    """Priekšmeta nosaukums un saknes mape — pēc noklusējuma dabaszinības."""
    return (at.get("prieksmets") or C.PRIEKSMETS,
            os.path.join(C.SAKNE, at["sakne"]) if at.get("sakne")
            else C.SAKNES_MAPE)


def galvene(doc, at, apaksvirsraksts, prieksmets):
    C.para(doc, "%s  |  %s  |  %s" % (prieksmets, C.SKOLA, C.GADS),
           size=9, color=C.GREY, after=1)
    C.para(doc, "Ātrais tests Nr. %d.  %s" % (at["nr"], at["nosaukums"]),
           size=16, bold=True, color=C.NAVY, after=1)
    C.para(doc, apaksvirsraksts, size=13, bold=True, color=C.NAVY, after=4)


H_GALVENE = (C.h_para("x", 9, after=1) + C.h_para("x", 16, after=1)
             + C.h_para("x", 13, after=4))


# ------------------------------------------------------------- skolēna lapa
def atbilzu_rindas(varianti):
    """Atbilžu variantu rindkopas — vienā rindā, ja ietilpst, citādi divās.

    Tas pats, ko pd_common.testa_rindas(), tikai mazākā fontā un ar
    šaurāku atstarpi, lai lapā ietilptu visi jautājumi.
    """
    apz = ["%s) %s" % (C.BURTI[j], a) for j, a in enumerate(varianti)]
    visi = ATSTARPE.join(apz)
    if C.ietilpst(visi, JAUT_PT, C.TEKSTA_PLATUMS - 0.6):
        return [visi]
    return [ATSTARPE.join(apz[:2]), ATSTARPE.join(apz[2:])]


def zimet_jautajumus(doc, jautajumi):
    for i, (jaut, varianti, _) in enumerate(jautajumi, start=1):
        C.para(doc, "%d. %s" % (i, jaut), size=JAUT_PT, after=1)
        rindas = atbilzu_rindas(varianti)
        for j, r in enumerate(rindas):
            C.para(doc, r, size=JAUT_PT,
                   after=3 if j == len(rindas) - 1 else 0, left=0.6)


def h_jautajumi(jautajumi):
    h = 0.0
    for jaut, varianti, _ in jautajumi:
        h += C.h_para("99. " + jaut, JAUT_PT, after=1)
        rindas = atbilzu_rindas(varianti)
        for j, r in enumerate(rindas):
            h += C.h_para(r, JAUT_PT, after=3 if j == len(rindas) - 1 else 0,
                          left=0.6)
    return h


def _tabulas_bloki(n):
    return [list(range(s, min(s + AILES, n))) for s in range(0, n, AILES)]


def atbilzu_tabula(doc, jautajumi, atslega=False):
    """Tabula lapas apakšā: augšā jautājuma numurs, apakšā A B C D.

    Skolēna lapā apakšējā rindā ir visi četri burti, ko apvilkt; atbilžu
    lapā tur ir tikai pareizais burts.
    """
    n = len(jautajumi)
    C.para(doc, "ATBILŽU ATSLĒGA" if atslega else "ATBILŽU TABULA",
           size=10, bold=True, color=C.NAVY, before=4, after=2)
    if not atslega:
        C.para(doc, "Katram jautājumam apvelc vienu burtu! Labojums der tikai "
                    "tad, ja nepareizā atbilde ir skaidri nosvītrota.",
               size=9, italic=True, color=C.GREY, after=3)
    bloki = _tabulas_bloki(n)
    t = doc.add_table(rows=2 * len(bloki), cols=AILES)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    plat = C.TEKSTA_PLATUMS / AILES
    for b, bloks in enumerate(bloki):
        nr_rinda, atb_rinda = t.rows[2 * b], t.rows[2 * b + 1]
        for j in range(AILES):
            nr_rinda.cells[j].width = Cm(plat)
            atb_rinda.cells[j].width = Cm(plat)
            if j >= len(bloks):
                C.cell_text(nr_rinda.cells[j], "", size=10)
                C.cell_text(atb_rinda.cells[j], "", size=12)
                continue
            i = bloks[j]
            C.shade(nr_rinda.cells[j], "1F3864")
            C.cell_text(nr_rinda.cells[j], "%d." % (i + 1), size=10,
                        bold=True, color=C.WHITE,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
            teksts = (C.BURTI[jautajumi[i][2]] if atslega
                      else BURTU_ATSTARPE.join(C.BURTI))
            C.cell_text(atb_rinda.cells[j], teksts, size=12, bold=atslega,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
        atb_rinda.height = Cm(0.70)
        atb_rinda.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    C.para(doc, after=0)


def h_atbilzu_tabula(n, atslega=False):
    h = C.h_para("x", 10, before=4, after=2)
    if not atslega:
        h += C.h_para("x" * 100, 9, after=3)
    return h + len(_tabulas_bloki(n)) * (0.55 + 0.70) + 0.45


def punktu_rinda(doc, kopa):
    C.para(doc, "Punkti:  ............ / %d          Balle:  ............     "
                "     Skolotāja paraksts:  ......................." % kopa,
           size=10, before=4, after=0)


H_PUNKTI = C.h_para("x" * 100, 10, before=4, after=0)


def build_darba_lapa(at, path, prieksmets):
    doc = C.new_doc()
    n = len(at["jautajumi"])
    galvene(doc, at, "%d jautājumi  ·  %d minūtes  ·  %d punkti"
            % (n, at["laiks"], n), prieksmets)
    ievads = ("Katram jautājumam ir tikai viena pareizā atbilde. Par pareizu "
              "atbildi — 1 punkts, par nepareizu punktus neatņem. Atļauts "
              "lietot kalkulatoru.")
    C.para(doc, ievads, size=9.5, italic=True, color=C.GREY, after=4)
    C.veidlapa(doc)
    C.atgadnes_kaste(doc, at["atgadne"])
    zimet_jautajumus(doc, at["jautajumi"])
    atbilzu_tabula(doc, at["jautajumi"])
    punktu_rinda(doc, n)
    doc.save(path)
    return path


def augstums(at):
    """Skolēna lapas novērtētais augstums centimetros."""
    return (H_GALVENE + C.h_para("x" * 175, 9.5, after=4) + 1.45
            + 0.55 + len(at["atgadne"]) * 0.42 + 0.45
            + h_jautajumi(at["jautajumi"])
            + h_atbilzu_tabula(len(at["jautajumi"]))
            + H_PUNKTI)


# --------------------------------------------------------- atbildes skolotājam
def build_atbildes(at, path, prieksmets):
    doc = C.new_doc()
    n = len(at["jautajumi"])
    C.para(doc, "%s  |  %s  |  %s" % (prieksmets, C.SKOLA, C.GADS),
           size=9, color=C.GREY, after=1)
    C.para(doc, "Ātrais tests Nr. %d.  %s" % (at["nr"], at["nosaukums"]),
           size=16, bold=True, color=C.NAVY, after=1)
    C.para(doc, "Atbildes un vērtēšanas skala — SKOLOTĀJAM", size=13,
           bold=True, color=C.NAVY, after=4)
    C.para(doc, "%s  Vērtē %s stundā apgūto. Izpildes laiks — %d minūtes, "
                "kopā %d punkti (1 p. par katru pareizu atbildi). Plānotais "
                "datums — %s"
           % (at["apraksts"], at["stundas"], at["laiks"], n, at["datums"]),
           size=10.5, color=C.GREY, after=4)

    atbilzu_tabula(doc, at["jautajumi"], atslega=True)

    C.para(doc, "Pareizās atbildes", size=12, bold=True, color=C.NAVY,
           before=6, after=2)
    for i, (_, varianti, pareizais) in enumerate(at["jautajumi"], start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        r1 = p.add_run("%2d. — %s) " % (i, C.BURTI[pareizais]))
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.name = "Calibri"
        C.write_runs(p, varianti[pareizais], 10, color=C.GREY)

    C.para(doc, "Ko tests pārbauda", size=12, bold=True, color=C.NAVY,
           before=8, after=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(["Jaut.", "Pārbaudāmais sasniedzamais rezultāts",
                           "Stunda"]):
        C.shade(t.rows[0].cells[i], "1F3864")
        C.cell_text(t.rows[0].cells[i], h, size=9.5, bold=True, color=C.WHITE)
    for rinda in at["struktura"]:
        cells = t.add_row().cells
        C.cell_text(cells[0], rinda[0], size=10, bold=True, color=C.NAVY)
        C.cell_text(cells[1], rinda[1], size=10)
        C.cell_text(cells[2], rinda[2], size=10, color=C.GREY)
    for i, w in enumerate([2.2, 13.2, 2.6]):
        for r in t.rows:
            r.cells[i].width = Cm(w)

    C.para(doc, "Vērtēšanas skala (kopā %d punkti)" % n, size=12, bold=True,
           color=C.NAVY, before=8, after=2)
    sk = C.skala(n)
    t = doc.add_table(rows=2, cols=len(sk))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (balle, no_p, lidz_p) in enumerate(sk):
        C.shade(t.rows[0].cells[i], "1F3864")
        C.cell_text(t.rows[0].cells[i], str(balle), size=10, bold=True,
                    color=C.WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        teksts = str(no_p) if no_p == lidz_p else "%d–%d" % (no_p, lidz_p)
        C.cell_text(t.rows[1].cells[i], teksts, size=10,
                    align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in t.rows:
            r.cells[i].width = Cm(C.TEKSTA_PLATUMS / len(sk))
    C.para(doc, "Augšējā rindā — balle, apakšējā — iegūto punktu skaits.",
           size=9, italic=True, color=C.GREY, before=2, after=0)

    doc.save(path)
    return path


# ----------------------------------------------------------------- pārbaude
def parbaudi(at):
    """Pārbauda atbilžu skaitu, izkliedi un to, vai lapa paliek viena."""
    jaut = at["jautajumi"]
    assert jaut, "%s: nav neviena jautājuma" % at["fails"]
    assert len(jaut) % AILES == 0, \
        "%s: jautājumu skaits (%d) nedalās ar %d — atbilžu tabula būs " \
        "nepilna" % (at["fails"], len(jaut), AILES)
    for teksts, varianti, pareizais in jaut:
        assert len(varianti) == 4, "%s: %s — vajag 4 atbilžu variantus" % (
            at["fails"], teksts[:40])
        assert 0 <= pareizais < 4
        assert len(set(varianti)) == 4, \
            "%s: %s — atkārtojas atbilžu varianti" % (at["fails"], teksts[:40])
    skaits = [0, 0, 0, 0]
    for _, _, pareizais in jaut:
        skaits[pareizais] += 1
    assert max(skaits) <= len(jaut) // 2, \
        "%s: pareizo atbilžu izkliede %s" % (at["fails"], skaits)
    assert min(skaits) >= 1, \
        "%s: kāds burts netiek lietots — izkliede %s" % (at["fails"], skaits)
    h = augstums(at)
    assert h <= C.LAPAS_BUDZETS, \
        "%s: lapa neietilpst — %.1f cm no %.1f cm" % (at["fails"], h,
                                                      C.LAPAS_BUDZETS)


def build(at):
    parbaudi(at)
    prieksmets, sakne = _konteksts(at)
    mape = os.path.join(sakne, at["mape"])
    if not os.path.isdir(mape):
        raise SystemExit("Nav atrasta mape: %s" % mape)
    izveidoti = [
        build_darba_lapa(at, os.path.join(mape, "%s.docx" % at["fails"]),
                         prieksmets),
        build_atbildes(at, os.path.join(mape, "%s (atbildes).docx"
                                        % at["fails"]), prieksmets),
    ]
    for c in izveidoti:
        print("Izveidots: %s  (lapa %.1f cm no %.1f cm)"
              % (os.path.relpath(c, C.SAKNE), augstums(at), C.LAPAS_BUDZETS))
    return izveidoti
