# -*- coding: utf-8 -*-
"""Aptuveni pārbauda, vai teksts neizplūst ārpus slaida / tekstlodziņa.

Novērtē rindu skaitu pēc rakstzīmju platuma un salīdzina ar tekstlodziņa
augstumu. Nav precīzs renderētājs, bet pietiekami, lai pamanītu pārpildi.
"""
import sys
from pptx import Presentation
from pptx.util import Pt

sys.stdout.reconfigure(encoding="utf-8")

EMU_IN = 914400
CHAR_W = 0.47      # vidējais rakstzīmes platums attiecībā pret fonta izmēru
LINE_H = 1.22      # rindstarpa


def analyse(path):
    prs = Presentation(path)
    SW = prs.slide_width / EMU_IN
    SH = prs.slide_height / EMU_IN
    problems = []
    for i, slide in enumerate(prs.slides, 1):
        # tekstlodziņi nedrīkst pārklāties savā starpā
        tbs = [s for s in slide.shapes
               if s.has_text_frame and s.text_frame.text.strip()
               and s.shape_type is None]
        for a_i in range(len(tbs)):
            for b_i in range(a_i + 1, len(tbs)):
                a, b = tbs[a_i], tbs[b_i]
                ox = min(a.left + a.width, b.left + b.width) - \
                    max(a.left, b.left)
                oy = min(a.top + a.height, b.top + b.height) - \
                    max(a.top, b.top)
                if ox > 0.05 * EMU_IN and oy > 0.05 * EMU_IN:
                    problems.append(
                        (i, "TEKSTI PĀRKLĀJAS (%.2f x %.2f collas)"
                            % (ox / EMU_IN, oy / EMU_IN),
                         a.text_frame.text.strip()[:34].replace("\n", " ")
                         + "  ><  "
                         + b.text_frame.text.strip()[:34].replace("\n", " ")))

        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            tf = shp.text_frame
            text = tf.text.strip()
            if not text:
                continue
            # Formulu gabali (vertikālās daļas) zīmēti bez aplaušanas un
            # apzināti šauri - tos rindu skaita pārbaude nav jāpiemēro.
            if tf.word_wrap is False:
                continue
            w_in = shp.width / EMU_IN
            h_in = shp.height / EMU_IN
            x_in = shp.left / EMU_IN
            y_in = shp.top / EMU_IN

            ml = tf.margin_left / EMU_IN
            mr = tf.margin_right / EMU_IN
            mt = tf.margin_top / EMU_IN
            mb = tf.margin_bottom / EMU_IN
            avail_w_pt = (w_in - ml - mr) * 72
            avail_h_pt = (h_in - mt - mb) * 72

            total = 0.0
            for p in tf.paragraphs:
                ptext = "".join(r.text for r in p.runs)
                sizes = [r.font.size.pt for r in p.runs
                         if r.font.size is not None]
                sz = max(sizes) if sizes else 18.0
                sb = p.space_before.pt if p.space_before is not None else 0
                sa = p.space_after.pt if p.space_after is not None else 0
                if not ptext:
                    total += sz * LINE_H + sb + sa
                    continue
                cpl = max(1, int(avail_w_pt / (sz * CHAR_W)))
                nlines = max(1, -(-len(ptext) // cpl))
                total += nlines * sz * LINE_H + sb + sa

            if total > avail_h_pt + 1:
                problems.append(
                    (i, "PĀRPILDE %.0f pt > %.0f pt" % (total, avail_h_pt),
                     text[:70].replace("\n", " ")))
            if x_in < -0.02 or y_in < -0.02 or \
               x_in + w_in > SW + 0.02 or y_in + h_in > SH + 0.02:
                problems.append(
                    (i, "ĀRPUS SLAIDA (%.2f, %.2f, %.2f x %.2f)"
                        % (x_in, y_in, w_in, h_in),
                     text[:70].replace("\n", " ")))
    return len(prs.slides._sldIdLst), problems


if __name__ == "__main__":
    n, probs = analyse(sys.argv[1])
    print("Slaidi: %d" % n)
    if not probs:
        print("OK - pārpildes nav konstatētas.")
    else:
        print("Atrastas %d iespējamās problēmas:" % len(probs))
        for sl, kind, t in probs:
            print("  slaids %2d | %s | %s" % (sl, kind, t))
