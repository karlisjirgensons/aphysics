# -*- coding: utf-8 -*-
"""
Kopīgais pārbaudes darbu (PD) veidotājs dabaszinību fizikas daļai.

No viena satura apraksta (skat. pd_XX_saturs.py) uzbūvē divus .docx failus:
  «<nosaukums>.docx»             — skolēna darba lapa ar visiem variantiem
  «<nosaukums> (atbildes).docx»  — atbildes, kritēriji un vērtēšanas skala

Lapas dalījums nav ierakstīts saturā — modulis novērtē katra bloka augstumu
un ieliek lappuses pārtraukumu tur, kur bloks vairs neietilptu.

Satura shēma
------------
PD = {
    "nr": 2, "nosaukums": ..., "mape": ..., "fails": ...,
    "stundas": "3.1.–3.12.", "datums": "16.10.2026.",
    "kopa": 30, "laiks": 40, "apraksts": ..., "atgadne": [...],
    "struktura": [(uzd, sasniedzamais rezultāts, stunda, punkti), ...],
    "varianti": [{"nr": "1. variants", "tests": [...], "uzdevumi": [...]}, ...],
}

Uzdevumu tipi
  parveide   — aizpildāmu rindu tabula:   rindas=[(jautājums, atbilde), ...]
  aprekins   — pilns risinājuma pieraksts: teksts, vieta, risinajums, kriteriji
  jautajumi  — vairāki apakšjautājumi:     ievads, jaut=[(teksts, p), ...],
                                           vieta, atbildes
"""

import os

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SKOLA = "Ādažu vidusskola"
GADS = "2026./2027. m. g."

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x60, 0x60, 0x60)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BURTI = "ABCD"
INDEKSI = "abcdefgh"

LAPAS_AUGSTUMS = 27.3        # cm — A4 mīnus piemales
LAPAS_BUDZETS = 26.2         # cm — cik daudz no tās aizpilda (rezerve novērtējuma kļūdai)
TEKSTA_PLATUMS = 18.0        # cm

SAKNE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DABASZINIBAS = os.path.join(SAKNE, "Dabaszinibas")

# Priekšmeta nosaukums un darbu saknes mape - fizikas kurss tos nomaina
# (sk. fiz_pd_common.py); dabaszinību darbiem vērtības nemainās.
PRIEKSMETS = "Dabaszinības (fizikas daļa)"
SAKNES_MAPE = DABASZINIBAS

# Vērtēšanas skala procentos no maksimālā punktu skaita.
SKALAS_ROBEZAS = [(10, 0.95), (9, 0.86), (8, 0.76), (7, 0.66), (6, 0.53),
                  (5, 0.43), (4, 0.30), (3, 0.20), (2, 0.10), (1, 0.0)]


# ------------------------------------------------------------- augstuma vērtējums
def rindu_skaits(teksts, izmers_pt, platums_cm):
    """Aptuvenais rindu skaits Calibri tekstam dotajā platumā."""
    if not teksts:
        return 1
    garums_cm = len(teksts) * izmers_pt * 0.47 / 28.35
    n = int(garums_cm / platums_cm)
    return max(1, n + (1 if garums_cm % platums_cm else 0))


def h_para(teksts, izmers_pt=11.0, before=0, after=3, left=0.0):
    n = rindu_skaits(teksts, izmers_pt, TEKSTA_PLATUMS - left)
    return n * izmers_pt * 1.25 / 28.35 + (before + after) / 28.35


def ietilpst(teksts, izmers_pt, platums_cm):
    return len(teksts) * izmers_pt * 0.47 / 28.35 <= platums_cm


# ----------------------------------------------------------------- pamatelementi
def para(doc, text="", size=11, bold=False, italic=False, color=None,
         before=0, after=3, left=0, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.left_indent = Cm(left)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = "Calibri"
        r.font.color.rgb = color or BLACK
    return p


def shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def cell_text(cell, text, size=10.5, bold=False, italic=False, color=None,
              align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Calibri"
    r.font.color.rgb = color or BLACK
    return p


def darba_vieta(doc, augstums, teksts=None):
    """Ierāmēts laukums, kurā skolēns raksta risinājumu."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    row = t.rows[0]
    row.height = Cm(augstums)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    c = row.cells[0]
    c.width = Cm(TEKSTA_PLATUMS)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if teksts:
        r = p.add_run(teksts)
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.name = "Calibri"
        r.font.color.rgb = GREY
    return t


def lapas_partraukums(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.5)
    s.top_margin = Cm(1.2)
    s.bottom_margin = Cm(1.2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(3)
    return doc


# ------------------------------------------------------------------- galvenes
def galvene(doc, pd, apaksvirsraksts):
    para(doc, "%s  |  %s  |  %s" % (PRIEKSMETS, SKOLA, GADS),
         size=9, color=GREY, after=1)
    para(doc, "Pārbaudes darbs Nr. %d.  %s" % (pd["nr"], pd["nosaukums"]),
         size=16, bold=True, color=NAVY, after=1)
    para(doc, apaksvirsraksts, size=13, bold=True, color=NAVY, after=4)


H_GALVENE = (h_para("x", 9, after=1) + h_para("x", 16, after=1)
             + h_para("x", 13, after=4))


def veidlapa(doc):
    t = doc.add_table(rows=2, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, a in enumerate(["Vārds", "Uzvārds", "Klase", "Datums"]):
        cell_text(t.rows[0].cells[i], "." * 32, size=11)
        cell_text(t.rows[1].cells[i], a, size=9, color=GREY)
    for i, w in enumerate([5.5, 5.5, 3.5, 3.5]):
        for r in t.rows:
            r.cells[i].width = Cm(w)
    para(doc, after=2)


def atgadnes_kaste(doc, rindas):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    c.width = Cm(TEKSTA_PLATUMS)
    shade(c, "EEF2F8")
    c.text = ""
    p0 = c.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r = p0.add_run("ATGĀDNE")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = NAVY
    for rinda in rindas:
        p = c.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(rinda)
        rr.font.size = Pt(9)
        rr.font.name = "Calibri"
    para(doc, after=2)


def uzdevuma_virsraksts(doc, nr, teksts, punkti, before=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("%d. uzdevums. %s " % (nr, teksts))
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = NAVY
    r2 = p.add_run("(%d p.)" % punkti)
    r2.font.size = Pt(11)
    r2.font.bold = True
    r2.font.name = "Calibri"
    r2.font.color.rgb = GREY
    return p


# --------------------------------------------------------- skolēna darba lapa
def testa_rindas(varianti):
    """Atgriež atbilžu rindkopu tekstus — vienā vai divās rindās."""
    apz = ["%s) %s" % (BURTI[j], a) for j, a in enumerate(varianti)]
    visi = "        ".join(apz)
    if ietilpst(visi, 10.5, TEKSTA_PLATUMS - 0.6):
        return [visi]
    return ["        ".join(apz[:2]), "        ".join(apz[2:])]


def h_tests(tests):
    h = h_para("x", 12, before=6, after=2) + h_para("x" * 80, 9.5, after=3)
    for jaut, varianti, _ in tests:
        h += h_para("99. " + jaut, 10.5, after=1)
        rindas = testa_rindas(varianti)
        for i, r in enumerate(rindas):
            h += h_para(r, 10.5, after=4 if i == len(rindas) - 1 else 0,
                        left=0.6)
    return h


def zimet_testu(doc, tests):
    uzdevuma_virsraksts(doc, 1, "Tests", len(tests), before=6)
    para(doc, "Katram jautājumam ir tikai viena pareizā atbilde. Apvelc "
              "pareizās atbildes burtu!", size=9.5, italic=True, color=GREY,
         after=3)
    for i, (jaut, varianti, _) in enumerate(tests, start=1):
        para(doc, "%d. %s" % (i, jaut), size=10.5, after=1)
        rindas = testa_rindas(varianti)
        for j, r in enumerate(rindas):
            para(doc, r, size=10.5, after=4 if j == len(rindas) - 1 else 0,
                 left=0.6)


def h_uzdevums(u):
    h = h_para("x", 12, before=8, after=2)
    if u["tips"] == "parveide":
        h += h_para("x" * 70, 9.5, after=3)
        h += len(u["rindas"]) * max(0.85, h_para("x", 11))
    elif u["tips"] == "aprekins":
        h += h_para(u["teksts"], 11, after=3) + u["vieta"]
    else:
        h += h_para(u["ievads"], 11, after=2)
        for jaut, p in u["jaut"]:
            h += h_para("9) %s  (%d p.)" % (jaut, p), 10.5, after=1, left=0.4)
        h += u["vieta"]
    return h


def zimet_uzdevumu(doc, nr, u):
    uzdevuma_virsraksts(doc, nr, u["virs"], u["punkti"])
    if u["tips"] == "parveide":
        para(doc, u.get("note", "Ieraksti trūkstošo vērtību! Par katru "
                                "pareizu pārveidojumu — 1 punkts."),
             size=9.5, italic=True, color=GREY, after=3)
        t = doc.add_table(rows=len(u["rindas"]), cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, (jaut, _) in enumerate(u["rindas"]):
            row = t.rows[i]
            row.height = Cm(0.85)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            row.cells[0].width = Cm(TEKSTA_PLATUMS)
            cell_text(row.cells[0], "%s)  %s" % (INDEKSI[i], jaut), size=11)
    elif u["tips"] == "aprekins":
        para(doc, u["teksts"], size=11, after=3)
        darba_vieta(doc, u["vieta"],
                    "Dots:                     Jāaprēķina:              "
                    "Formulas:                Aprēķins:                "
                    "Atbilde:")
    else:
        para(doc, u["ievads"], size=11, after=2)
        for i, (jaut, p) in enumerate(u["jaut"], start=1):
            para(doc, "%d) %s  (%d p.)" % (i, jaut, p), size=10.5, after=1,
                 left=0.4)
        darba_vieta(doc, u["vieta"])


def vertejuma_aile(doc, pd, uzdevumi):
    n = len(uzdevumi) + 1
    galva = ["%d. uzd." % i for i in range(1, n + 1)] + ["Kopā", "Balle"]
    t = doc.add_table(rows=2, cols=len(galva))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    plat = TEKSTA_PLATUMS / len(galva)
    for i, g in enumerate(galva):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], g, size=9, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t.rows[1].cells[i], "", size=11)
        for r in t.rows:
            r.cells[i].width = Cm(plat)
    t.rows[1].height = Cm(1.0)
    t.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    punkti = [len(pd["varianti"][0]["tests"])] + [u["punkti"] for u in uzdevumi]
    para(doc, "Maksimālais punktu skaits par uzdevumu:  %s"
         % "  ·  ".join("%d. uzd. — %d p." % (i + 1, p)
                        for i, p in enumerate(punkti)),
         size=8.5, color=GREY, before=2)


H_VERTEJUMS = 1.1 + h_para("x" * 110, 8.5, before=2)


def variants_lapa(doc, pd, v):
    """Uzzīmē vienu variantu, pats izlemjot, kur beidzas lapa."""
    galvene(doc, pd, v["nr"])
    ievads = ("Darba izpildes laiks — %d minūtes. Maksimālais punktu "
              "skaits — %d punkti. Atļauts lietot kalkulatoru un formulu "
              "lapu. Aprēķinu uzdevumos jāparāda risinājuma gaita: "
              "Dots → Jāaprēķina → Formulas → Aprēķins → Atbilde."
              % (pd["laiks"], pd["kopa"]))
    para(doc, ievads, size=9.5, italic=True, color=GREY, after=4)
    veidlapa(doc)
    atgadnes_kaste(doc, pd["atgadne"])

    augstums = (H_GALVENE + h_para(ievads, 9.5, after=4) + 1.45
                + 0.55 + len(pd["atgadne"]) * 0.42 + 0.45)

    bloki = [(h_tests(v["tests"]), lambda d: zimet_testu(d, v["tests"]))]
    for i, u in enumerate(v["uzdevumi"], start=2):
        bloki.append((h_uzdevums(u),
                      lambda d, n=i, uu=u: zimet_uzdevumu(d, n, uu)))
    bloki.append((H_VERTEJUMS,
                  lambda d: vertejuma_aile(d, pd, v["uzdevumi"])))

    for h, zimet in bloki:
        if augstums + h > LAPAS_BUDZETS:
            lapas_partraukums(doc)
            galvene(doc, pd, v["nr"])
            augstums = H_GALVENE
        zimet(doc)
        augstums += h


def build_darba_lapa(pd, path):
    doc = new_doc()
    for i, v in enumerate(pd["varianti"]):
        if i:
            lapas_partraukums(doc)
        variants_lapa(doc, pd, v)
    doc.save(path)
    return path


# ------------------------------------------------------- atbildes skolotājam
def skala(kopa):
    """Punktu robežas ballēm 10…1 no procentu robežām."""
    out, augsa = [], kopa
    for balle, dala in SKALAS_ROBEZAS:
        apaksa = 0 if balle == 1 else int(round(dala * kopa))
        apaksa = min(apaksa, augsa)
        out.append((balle, apaksa, augsa))
        augsa = apaksa - 1
    return out


def zimet_testa_atbildes(doc, tests):
    para(doc, "1. uzdevums. Tests (%d p. — par katru pareizu atbildi 1 p.)"
         % len(tests), size=12, bold=True, color=NAVY, before=4, after=2)
    for i, (_, varianti, pareizais) in enumerate(tests, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        r1 = p.add_run("%2d. — %s) " % (i, BURTI[pareizais]))
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.name = "Calibri"
        r2 = p.add_run(varianti[pareizais])
        r2.font.size = Pt(10)
        r2.font.name = "Calibri"
        r2.font.color.rgb = GREY


def h_testa_atbildes(tests):
    h = h_para("x", 12, before=4, after=2)
    for _, varianti, pareizais in tests:
        h += h_para("99. — A) " + varianti[pareizais], 10.5, after=0, left=0.4)
    return h


def zimet_uzd_atbildes(doc, nr, u):
    para(doc, "%d. uzdevums. %s (%d p.)" % (nr, u["virs"], u["punkti"]),
         size=12, bold=True, color=NAVY, before=8, after=2)
    if u["tips"] == "parveide":
        for i, (jaut, atb) in enumerate(u["rindas"]):
            tirs = jaut.replace("." * 21, "___")
            para(doc, "%s)  %s   →   %s" % (INDEKSI[i], tirs, atb),
                 size=10, after=0, left=0.4)
        para(doc, u.get("kriterijs",
                        "Punktu nepiešķir, ja trūkst mērvienības vai "
                        "skaitlis nav standartformā tur, kur tas prasīts."),
             size=9, italic=True, color=GREY, before=2, after=0)
    elif u["tips"] == "aprekins":
        for rinda in u["risinajums"]:
            para(doc, rinda, size=10, after=0, left=0.4)
        para(doc, "Vērtēšana:", size=10, bold=True, before=3, after=1,
             left=0.4)
        for k in u["kriteriji"]:
            para(doc, "•  " + k, size=9.5, after=0, left=0.8)
    else:
        for rinda in u["atbildes"]:
            para(doc, rinda, size=10, after=2, left=0.4)


def h_uzd_atbildes(u):
    h = h_para("x", 12, before=8, after=2)
    if u["tips"] == "parveide":
        for i, (jaut, atb) in enumerate(u["rindas"]):
            h += h_para("%s)  %s   →   %s" % (INDEKSI[i], jaut, atb), 10,
                        after=0, left=0.4)
        h += h_para("x" * 100, 9, before=2, after=0)
    elif u["tips"] == "aprekins":
        for rinda in u["risinajums"]:
            h += h_para(rinda, 10, after=0, left=0.4)
        h += h_para("x", 10, before=3, after=1, left=0.4)
        for k in u["kriteriji"]:
            h += h_para("•  " + k, 9.5, after=0, left=0.8)
    else:
        for rinda in u["atbildes"]:
            h += h_para(rinda, 10, after=2, left=0.4)
    return h


def atbilzu_lapa(doc, pd, v):
    virsraksts = "%s — atbildes un vērtēšanas kritēriji" % v["nr"]
    para(doc, virsraksts, size=14, bold=True, color=NAVY, before=6, after=4)
    augstums = h_para(virsraksts, 14, before=6, after=4)

    bloki = [(h_testa_atbildes(v["tests"]),
              lambda d: zimet_testa_atbildes(d, v["tests"]))]
    for i, u in enumerate(v["uzdevumi"], start=2):
        bloki.append((h_uzd_atbildes(u),
                      lambda d, n=i, uu=u: zimet_uzd_atbildes(d, n, uu)))

    for h, zimet in bloki:
        if augstums + h > LAPAS_BUDZETS:
            lapas_partraukums(doc)
            para(doc, "%s (turpinājums)" % v["nr"], size=12, bold=True,
                 color=NAVY, after=4)
            augstums = h_para("x", 12, after=4)
        zimet(doc)
        augstums += h


def build_atbildes(pd, path):
    doc = new_doc()
    para(doc, "%s  |  %s  |  %s" % (PRIEKSMETS, SKOLA, GADS),
         size=9, color=GREY, after=1)
    para(doc, "Pārbaudes darbs Nr. %d.  %s" % (pd["nr"], pd["nosaukums"]),
         size=16, bold=True, color=NAVY, after=1)
    para(doc, "Atbildes, vērtēšanas kritēriji un skala — SKOLOTĀJAM",
         size=13, bold=True, color=NAVY, after=4)
    para(doc, "%s  Vērtē %s stundā apgūto. Darba izpildes laiks — %d "
              "minūtes, kopā %d punkti. Plānotais datums — %s"
         % (pd["apraksts"], pd["stundas"], pd["laiks"], pd["kopa"],
            pd["datums"]), size=10.5, color=GREY, after=4)

    para(doc, "Darba struktūra", size=12, bold=True, color=NAVY, before=4,
         after=2)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(["Uzd.", "Pārbaudāmais sasniedzamais rezultāts",
                           "Stunda", "Punkti"]):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], h, size=9.5, bold=True, color=WHITE)
    for rinda in pd["struktura"]:
        cells = t.add_row().cells
        cell_text(cells[0], rinda[0], size=10, bold=True, color=NAVY)
        cell_text(cells[1], rinda[1], size=10)
        cell_text(cells[2], rinda[2], size=10, color=GREY)
        cell_text(cells[3], str(rinda[3]), size=10, bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, w in enumerate([1.4, 12.2, 2.4, 1.8]):
        for r in t.rows:
            r.cells[i].width = Cm(w)

    para(doc, "Vērtēšanas skala (kopā %d punkti)" % pd["kopa"], size=12,
         bold=True, color=NAVY, before=8, after=2)
    sk = skala(pd["kopa"])
    t = doc.add_table(rows=2, cols=len(sk))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (balle, no_p, lidz_p) in enumerate(sk):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], str(balle), size=10, bold=True,
                  color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        teksts = str(no_p) if no_p == lidz_p else "%d–%d" % (no_p, lidz_p)
        cell_text(t.rows[1].cells[i], teksts, size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        for r in t.rows:
            r.cells[i].width = Cm(TEKSTA_PLATUMS / len(sk))
    para(doc, "Augšējā rindā — balle, apakšējā — iegūto punktu skaits.",
         size=9, italic=True, color=GREY, before=2, after=4)

    for v in pd["varianti"]:
        lapas_partraukums(doc)
        atbilzu_lapa(doc, pd, v)

    doc.save(path)
    return path


# ------------------------------------------------------------------- pārbaude
def parbaudi(pd):
    """Pārbauda punktu summu un atbilžu izkliedi; kļūdu gadījumā met AssertionError."""
    for v in pd["varianti"]:
        kopa = len(v["tests"]) + sum(u["punkti"] for u in v["uzdevumi"])
        assert kopa == pd["kopa"], \
            "%s %s: punktu summa %d, gaidīts %d" % (pd["fails"], v["nr"],
                                                    kopa, pd["kopa"])
        for u in v["uzdevumi"]:
            if u["tips"] == "parveide":
                assert len(u["rindas"]) == u["punkti"], \
                    "%s %s: %s — rindu skaits neatbilst punktiem" % (
                        pd["fails"], v["nr"], u["virs"])
            elif u["tips"] == "jautajumi":
                assert sum(p for _, p in u["jaut"]) == u["punkti"], \
                    "%s %s: %s — apakšpunktu summa neatbilst" % (
                        pd["fails"], v["nr"], u["virs"])
        for jaut, varianti, pareizais in v["tests"]:
            assert len(varianti) == 4, "%s: %s — vajag 4 atbilžu variantus" % (
                pd["fails"], jaut[:40])
            assert 0 <= pareizais < 4
        # neviens burts nedrīkst atkārtoties vairāk par pusi jautājumu
        skaits = [0, 0, 0, 0]
        for _, _, pareizais in v["tests"]:
            skaits[pareizais] += 1
        assert max(skaits) <= len(v["tests"]) // 2, \
            "%s %s: pareizo atbilžu izkliede %s" % (pd["fails"], v["nr"],
                                                    skaits)


def build(pd):
    """Uzbūvē abus failus temata mapē (un, ja norādīts, arī papildu mapēs).

    Atslēga "papildu_mapes" der apvienotiem darbiem: PD6 vērtē gan 8., gan
    9. tematu, tāpēc tas pats fails jāatrod abu tematu mapēs. Abas kopijas
    veido viens un tas pats saturs, tāpēc tās nevar atšķirties.
    """
    parbaudi(pd)
    izveidoti = []
    for i, mapes_nos in enumerate([pd["mape"]] + list(pd.get("papildu_mapes",
                                                             []))):
        mape = os.path.join(SAKNES_MAPE, mapes_nos)
        if not os.path.isdir(mape):
            raise SystemExit("Nav atrasta mape: %s" % mape)
        izveidoti.append(
            build_darba_lapa(pd, os.path.join(mape, "%s.docx" % pd["fails"])))
        izveidoti.append(
            build_atbildes(pd, os.path.join(
                mape, "%s (atbildes).docx" % pd["fails"])))
    for c in izveidoti:
        print("Izveidots: %s" % os.path.relpath(c, SAKNE))
    return izveidoti
