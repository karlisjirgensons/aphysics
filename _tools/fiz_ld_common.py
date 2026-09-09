# -*- coding: utf-8 -*-
"""
Fizika I laboratorijas darbu (LD) veidotājs.

No viena satura apraksta (skat. fizld10_N.py / fizld11_N.py) uzbūvē divus
.docx failus temata mapē:

  «LDn. Nosaukums_tt.docx»                 — skolēna protokola veidne
  «LDn. Nosaukums_tt (skolotājam).docx»    — gaidāmie rezultāti, atbildes,
                                             vērtēšanas kritēriji un skala

Lapas dalījums nav ierakstīts saturā — modulis novērtē katra bloka augstumu
un ieliek lappuses pārtraukumu tur, kur bloks vairs neietilptu (tāpat kā
pd_common.py).

Satura shēma
------------
LD = {
    "nr": 2, "klase": "10. klase", "nosaukums": ..., "mape": ..., "fails": ...,
    "datums": "02.12.2026.", "svars": 9, "laiks": 80,
    "jautajums": "pētāmais jautājums",
    "merkis": ..., "hipoteze": "norāde, ko skolēns pieraksta",
    "teorija": [rindas atgādnei],
    "piederumi": [...],
    "drosiba": [...],
    "gaita": [(solis, apraksts), ...],
    "tabula": {"galva": [...], "rindas": 6, "platumi": [...]},
    "apstrade": [(virsraksts, norāde, vieta_cm), ...],
    "jautajumi": [(teksts, vieta_cm), ...],
    "kriteriji": [(kritērijs, punkti), ...],
    "kopa": 20,
    "gaidamie": [rindas skolotājam],
    "atbildes": [rindas skolotājam],
}

Palaiž:  .venv/Scripts/python.exe _tools/gen_fiz_ld.py
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                        # noqa: E402
from docx.shared import Cm, Pt                                       # noqa: E402

import pd_common as C                                                # noqa: E402
from pd_common import (GREY, NAVY, WHITE, TEKSTA_PLATUMS,            # noqa: E402
                       LAPAS_BUDZETS, cell_text, darba_vieta, h_para,
                       lapas_partraukums, new_doc, para, shade, skala)

SAKNE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIZIKA = os.path.join(SAKNE, "Fizika_1")
PRIEKSMETS = "Fizika I"


# ------------------------------------------------------------------ palīgi
class Lapa(object):
    """Seko līdzi aizpildītajam augstumam un ieliek lappuses pārtraukumus."""

    def __init__(self, doc):
        self.doc = doc
        self.h = 0.0

    def vieta(self, augstums):
        if self.h + augstums > LAPAS_BUDZETS:
            lapas_partraukums(self.doc)
            self.h = 0.0
        self.h += augstums

    def pieskaita(self, augstums):
        self.h += augstums


def galvene(doc, ld, apaksvirsraksts):
    para(doc, "%s  |  %s  |  %s  |  %s"
         % (PRIEKSMETS, ld["klase"], C.SKOLA, C.GADS),
         size=9, color=GREY, after=1)
    para(doc, "Laboratorijas darbs Nr. %d.  %s" % (ld["nr"], ld["nosaukums"]),
         size=16, bold=True, color=NAVY, after=1)
    para(doc, apaksvirsraksts, size=13, bold=True, color=NAVY, after=4)


def veidlapa(doc):
    t = doc.add_table(rows=2, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, a in enumerate(["Vārds, uzvārds", "Klase", "Grupas biedri",
                           "Datums"]):
        cell_text(t.rows[0].cells[i], "." * 30, size=11)
        cell_text(t.rows[1].cells[i], a, size=9, color=GREY)
    for i, w in enumerate([5.0, 2.5, 6.5, 4.0]):
        for r in t.rows:
            r.cells[i].width = Cm(w)
    para(doc, after=2)


def kaste(doc, virsraksts, rindas, fons="EEF2F8", size=9.5):
    """Ierāmēta, tonēta kaste ar virsrakstu un tekstu rindām."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    c = t.rows[0].cells[0]
    c.width = Cm(TEKSTA_PLATUMS)
    shade(c, fons)
    c.text = ""
    p0 = c.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r = p0.add_run(virsraksts)
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = NAVY
    for rinda in rindas:
        p = c.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(rinda)
        rr.font.size = Pt(size)
        rr.font.name = "Calibri"
    para(doc, after=2)


def h_kaste(virsraksts, rindas, size=9.5):
    h = h_para("x", 9, after=0) + 0.25
    for r in rindas:
        h += h_para(r, size, after=0)
    return h + h_para("", 11, after=2)


def sadalas_virsraksts(doc, nr, teksts, punkti=None, before=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("%s  %s " % (nr, teksts))
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = NAVY
    if punkti:
        r2 = p.add_run("(%d p.)" % punkti)
        r2.font.size = Pt(11)
        r2.font.bold = True
        r2.font.name = "Calibri"
        r2.font.color.rgb = GREY
    return p


def datu_tabula(doc, spec):
    galva = spec["galva"]
    n = spec["rindas"]
    plat = spec.get("platumi") or [TEKSTA_PLATUMS / len(galva)] * len(galva)
    t = doc.add_table(rows=n + 1, cols=len(galva))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, g in enumerate(galva):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], g, size=9.5, bold=True, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in range(1, n + 1):
        t.rows[r].height = Cm(0.72)
        t.rows[r].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i in range(len(galva)):
            cell_text(t.rows[r].cells[i], "", size=11)
    for i, w in enumerate(plat):
        for r in t.rows:
            r.cells[i].width = Cm(w)
    para(doc, after=2)
    return t


def h_tabula(spec):
    return (0.75 + spec["rindas"] * 0.75) + h_para("", 11, after=2)


# --------------------------------------------------- skolēna protokola lapa
def build_protokols(ld, path):
    doc = new_doc()
    L = Lapa(doc)

    galvene(doc, ld, "Skolēna protokols")
    L.pieskaita(h_para("x", 9, after=1) + h_para("x", 16, after=1)
                + h_para("x", 13, after=4))
    veidlapa(doc)
    L.pieskaita(1.65)

    info = ["Pētāmais jautājums:  %s" % ld["jautajums"],
            "Darba mērķis:  %s" % ld["merkis"],
            "Darba laiks:  %d min (trešdienas dubultstunda).   "
            "Vērtējums:  %d %% no gada atzīmes,  kopā %d punkti."
            % (ld["laiks"], ld["svars"], ld["kopa"]),
            "Protokolu iesniedz e-klasē nedēļas laikā pēc mērījumiem."]
    kaste(doc, "DARBA UZDEVUMS", info)
    L.pieskaita(h_kaste("DARBA UZDEVUMS", info))

    if ld.get("teorija"):
        kaste(doc, "ATGĀDNE", ld["teorija"])
        L.pieskaita(h_kaste("ATGĀDNE", ld["teorija"]))

    # 1. hipotēze
    L.vieta(h_para("x", 12, before=8, after=2) + h_para("x", 10, after=2)
            + 2.4)
    sadalas_virsraksts(doc, "1.", "Hipotēze", 2)
    para(doc, ld["hipoteze"], size=10, italic=True, color=GREY, after=2)
    darba_vieta(doc, 2.4)

    # 2. piederumi un drošība
    L.vieta(h_para("x", 12, before=8, after=2)
            + h_kaste("DARBA PIEDERUMI", ld["piederumi"])
            + h_kaste("DROŠĪBAS NOTEIKUMI", ld["drosiba"]))
    sadalas_virsraksts(doc, "2.", "Darba piederumi un drošība")
    kaste(doc, "DARBA PIEDERUMI", ld["piederumi"])
    kaste(doc, "DROŠĪBAS NOTEIKUMI", ld["drosiba"], fons="FBEEEE")

    # 3. darba gaita
    gaitas_h = h_para("x", 12, before=8, after=2)
    for i, (_, apr) in enumerate(ld["gaita"], 1):
        gaitas_h += h_para("%d) %s" % (i, apr), 10.5, after=2, left=0.4)
    L.vieta(gaitas_h)
    sadalas_virsraksts(doc, "3.", "Darba gaita")
    for i, (solis, apr) in enumerate(ld["gaita"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.4)
        r = p.add_run("%d) %s  " % (i, solis))
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.name = "Calibri"
        r2 = p.add_run(apr)
        r2.font.size = Pt(10.5)
        r2.font.name = "Calibri"

    # 4. mērījumu tabula
    L.vieta(h_para("x", 12, before=8, after=2) + h_para("x", 10, after=2)
            + h_tabula(ld["tabula"]))
    sadalas_virsraksts(doc, "4.", "Mērījumu dati", 5)
    para(doc, ld.get("tabulas_note",
                     "Datus raksti ar mērierīces precizitātei atbilstošu "
                     "ciparu skaitu un norādi mērvienības!"),
         size=10, italic=True, color=GREY, after=2)
    datu_tabula(doc, ld["tabula"])

    # 5. datu apstrāde
    for j, (virs, norade, vieta) in enumerate(ld["apstrade"]):
        h = h_para("x", 12, before=8 if j == 0 else 6, after=2)
        h += h_para(norade, 10.5, after=2) + vieta + 0.3
        L.vieta(h)
        if j == 0:
            sadalas_virsraksts(doc, "5.", "Datu apstrāde un aprēķini",
                               ld.get("apstrades_punkti", 5))
            para(doc, "Risinājumu pieraksti latviešu standartā: Dots → "
                      "Jāaprēķina → Formulas → Aprēķins → Atbilde.",
                 size=10, italic=True, color=GREY, after=2)
            L.pieskaita(h_para("x", 10, after=2))
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("%s.%d.  %s" % (5, j + 1, virs))
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.name = "Calibri"
        para(doc, norade, size=10.5, after=2)
        darba_vieta(doc, vieta)

    # 6. secinājumi un jautājumi
    h = h_para("x", 12, before=8, after=2)
    for teksts, vieta in ld["jautajumi"]:
        h += h_para(teksts, 10.5, after=2) + vieta + 0.3
    L.vieta(min(h, LAPAS_BUDZETS))
    sadalas_virsraksts(doc, "6.", "Secinājumi un izvērtējums",
                       ld.get("secinajumu_punkti", 6))
    for i, (teksts, vieta) in enumerate(ld["jautajumi"], 1):
        para(doc, "%d) %s" % (i, teksts), size=10.5, after=2, left=0.4)
        darba_vieta(doc, vieta)

    # vērtējuma aile
    L.vieta(h_para("x", 11, before=8, after=2) + 1.6)
    para(doc, "Vērtējums", size=11, bold=True, color=NAVY, before=8, after=2)
    galva = ["Hipotēze", "Mērījumi", "Aprēķini", "Secinājumi", "Noformējums",
             "Kopā", "Balle"]
    t = doc.add_table(rows=2, cols=len(galva))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, g in enumerate(galva):
        shade(t.rows[0].cells[i], "EEF2F8")
        cell_text(t.rows[0].cells[i], g, size=9, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        t.rows[1].cells[i].width = Cm(TEKSTA_PLATUMS / len(galva))
        t.rows[0].cells[i].width = Cm(TEKSTA_PLATUMS / len(galva))
    t.rows[1].height = Cm(1.0)
    t.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    doc.save(path)
    return path


# ------------------------------------------------------------ skolotāja lapa
def build_skolotajam(ld, path):
    doc = new_doc()
    galvene(doc, ld, "Skolotājam: gaidāmie rezultāti un vērtēšana")
    veidlapa_info = [
        "Norise:  %s, trešdienas dubultstunda (%d min)."
        % (ld["datums"], ld["laiks"]),
        "Vērtējums:  %d %% no gada atzīmes;  maksimums %d punkti."
        % (ld["svars"], ld["kopa"]),
        "Skolēni strādā %s; protokolu katrs iesniedz individuāli."
        % ld.get("darba_forma", "2-3 cilvēku grupās"),
    ]
    kaste(doc, "ORGANIZĀCIJA", veidlapa_info)

    para(doc, "Sagatavošana", size=12, bold=True, color=NAVY, before=8,
         after=2)
    for i, r in enumerate(ld["sagatavosana"], 1):
        para(doc, "%d) %s" % (i, r), size=10.5, after=2, left=0.4)

    para(doc, "Gaidāmie rezultāti", size=12, bold=True, color=NAVY,
         before=8, after=2)
    for r in ld["gaidamie"]:
        para(doc, "•  " + r, size=10.5, after=2, left=0.4)

    para(doc, "Atbildes uz secinājumu jautājumiem", size=12, bold=True,
         color=NAVY, before=8, after=2)
    for i, r in enumerate(ld["atbildes"], 1):
        para(doc, "%d) %s" % (i, r), size=10.5, after=2, left=0.4)

    lapas_partraukums(doc)

    para(doc, "Vērtēšanas kritēriji", size=12, bold=True, color=NAVY,
         before=0, after=3)
    t = doc.add_table(rows=len(ld["kriteriji"]) + 2, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, g in enumerate(["Nr.", "Kritērijs", "Punkti"]):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], g, size=9.5, bold=True, color=WHITE)
    for i, (kr, p) in enumerate(ld["kriteriji"], 1):
        cell_text(t.rows[i].cells[0], "%d." % i, size=10)
        cell_text(t.rows[i].cells[1], kr, size=10)
        cell_text(t.rows[i].cells[2], str(p), size=10,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    r = t.rows[len(ld["kriteriji"]) + 1]
    shade(r.cells[1], "EEF2F8")
    shade(r.cells[2], "EEF2F8")
    cell_text(r.cells[1], "KOPĀ", size=10, bold=True)
    cell_text(r.cells[2], str(ld["kopa"]), size=10, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, w in enumerate([1.2, 14.0, 2.8]):
        for row in t.rows:
            row.cells[i].width = Cm(w)
    para(doc, after=2)

    para(doc, "Vērtēšanas skala", size=12, bold=True, color=NAVY, before=8,
         after=3)
    robezas = skala(ld["kopa"])
    t2 = doc.add_table(rows=2, cols=len(robezas))
    t2.style = "Table Grid"
    t2.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (balle, apaksa, augsa) in enumerate(robezas):
        shade(t2.rows[0].cells[i], "EEF2F8")
        cell_text(t2.rows[0].cells[i], str(balle), size=9.5, bold=True,
                  color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(t2.rows[1].cells[i], "%d-%d" % (apaksa, augsa), size=9.5,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        for row in t2.rows:
            row.cells[i].width = Cm(TEKSTA_PLATUMS / len(robezas))
    para(doc, "Augšējā rindā - balle, apakšējā - punktu intervāls.",
         size=9, italic=True, color=GREY, after=2)

    if ld.get("piezimes"):
        para(doc, "Metodiskās piezīmes", size=12, bold=True, color=NAVY,
             before=8, after=2)
        for r in ld["piezimes"]:
            para(doc, "•  " + r, size=10.5, after=2, left=0.4)

    doc.save(path)
    return path


# ------------------------------------------------------------------ būvēšana
def parbaudi(ld):
    kopa = sum(p for _, p in ld["kriteriji"])
    assert kopa == ld["kopa"], \
        "%s: kritēriju summa %d, gaidīts %d" % (ld["fails"], kopa, ld["kopa"])
    assert len(ld["jautajumi"]) >= 3, "%s: par maz secinājumu jautājumu" % (
        ld["fails"])
    assert len(ld["atbildes"]) == len(ld["jautajumi"]), \
        "%s: atbilžu skaits neatbilst jautājumu skaitam" % ld["fails"]


def build(ld):
    parbaudi(ld)
    mape = os.path.join(FIZIKA, ld["mape"])
    if not os.path.isdir(mape):
        raise SystemExit("Nav atrasta mape: %s" % mape)
    a = build_protokols(ld, os.path.join(mape, "%s.docx" % ld["fails"]))
    b = build_skolotajam(ld, os.path.join(mape,
                                          "%s (skolotājam).docx"
                                          % ld["fails"]))
    for c in (a, b):
        print("Izveidots: %s" % os.path.relpath(c, SAKNE))
    return [a, b]
