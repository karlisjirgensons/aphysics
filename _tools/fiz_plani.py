# -*- coding: utf-8 -*-
"""
Fizika I - tematu, parbaudes darbu (PD) un laboratorijas darbu (LD) plani.

Paraugs: Fizika_1/theme_example.pdf. Atskiriba no dabaszinibu planiem seit
vertesanas kalendara ir svara (%) aile un darbu veidi PD, LD un PR.

Macibu kalendars (rules_fizika.txt):
  - macibas sakas 03.09.2026.;
  - fizikas stundas notiek TIKAI tresdienas un piektdienas;
  - tresdiena divas stundas pec kartas (dubultstunda), piektdiena viena -
    kopa tris stundas nedela;
  - laboratorijas darbi vienmer aiznem tresdienas dubultstundu.

Saturs:
  10. klase - temati 1.-5.
  11. klase - temati 6.-14.
"""

import datetime as dt
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SKOLA = "Ādažu vidusskola"
GADS = "2026./2027."
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x60, 0x60, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SAKNE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIZIKA = os.path.join(SAKNE, "Fizika_1")

# ------------------------------------------------------------ mācību kalendārs
SAKUMS = dt.date(2026, 9, 3)          # mācību gada sākums (ceturtdiena)
BEIGAS = dt.date(2027, 5, 31)

BRIVLAIKI = [(dt.date(2026, 10, 19), dt.date(2026, 10, 23)),
             (dt.date(2026, 12, 23), dt.date(2027, 1, 4)),
             (dt.date(2027, 3, 15), dt.date(2027, 3, 19))]
SVETKI = [dt.date(2026, 11, 18),      # Latvijas Republikas proklamēšanas diena
          dt.date(2027, 3, 26),       # Lielā Piektdiena
          dt.date(2027, 5, 4)]        # Neatkarības atjaunošanas diena

TRESDIENA, CETURTDIENA, PIEKTDIENA = 2, 3, 4


def d(x):
    return x.strftime("%d.%m.%Y")


class Grafiks:
    """Kad un kur notiek stundas - vienīgā vieta, kur zināms mācību ritms.

    Saturs (temati, PD, LD) no grafika nav atkarīgs, tāpēc vienu un to pašu
    plānu var izdrukāt vairākiem stundu sarakstiem - mainās tikai datumi.

    dienas   - {nedēļas diena: stundu skaits tajā dienā};
    ritms    - teikums dokumenta kalendāra daļai;
    bloka_vieta - kur nonāk dubultstundas (laboratorijas darbi).
    """

    def __init__(self, skola, dienas, ritms, bloka_vieta,
                 sakums=SAKUMS, beigas=BEIGAS,
                 brivlaiki=BRIVLAIKI, svetki=SVETKI):
        self.skola = skola
        self.dienas = dienas
        self.ritms = ritms
        self.bloka_vieta = bloka_vieta
        self.sakums, self.beigas = sakums, beigas
        self.brivlaiki, self.svetki = brivlaiki, svetki

    def macibu_dienas(self):
        """[(datums, stundu skaits tajā dienā), ...]"""
        out, cur = [], self.sakums
        while cur <= self.beigas:
            n = self.dienas.get(cur.weekday())
            if (n and not any(a <= cur <= b for a, b in self.brivlaiki)
                    and cur not in self.svetki):
                out.append((cur, n))
            cur += dt.timedelta(days=1)
        return out

    def slotu_saraksts(self):
        """Katrai mācību stundai (datums, vieta dienā)."""
        return [(datums, k)
                for datums, n in self.macibu_dienas()
                for k in range(n)]

    @staticmethod
    def _posms(a, b):
        """«19.-23.10.2026.», bet pāri mēnešiem «23.12.2026.-04.01.2027.»"""
        if (a.year, a.month) == (b.year, b.month):
            return "%d.-%s." % (a.day, d(b))
        return "%s.-%s." % (d(a), d(b))

    @staticmethod
    def _uzskaite(gabali):
        """«a, b un c» - pēdējo saista ar «un», nevis ar komatu."""
        gabali = list(gabali)
        if len(gabali) < 2:
            return "".join(gabali)
        return "%s un %s" % (", ".join(gabali[:-1]), gabali[-1])

    def apraksts(self):
        """Kalendāra rindkopa dokumentā - datumi un brīvdienas."""
        return ("Mācību gada sākums %s., noslēgums %s. Fizikas stundas "
                "notiek %s. Brīvlaiki: %s Stundu nav %s Datumi ir "
                "provizoriski un aprēķināti pēc šī kalendāra."
                % (d(self.sakums), d(self.beigas), self.ritms,
                   self._uzskaite(self._posms(a, b)
                                  for a, b in self.brivlaiki),
                   self._uzskaite("%s." % d(x) for x in self.svetki)))


ADAZI = Grafiks(
    SKOLA, {TRESDIENA: 2, PIEKTDIENA: 1},
    "TIKAI trešdienās (dubultstunda) un piektdienās (viena stunda)",
    "trešdienas dubultstundā")

CARNIKAVA = Grafiks(
    SKOLA, {CETURTDIENA: 3},
    "TIKAI ceturtdienās - trīs stundas pēc kārtas",
    "ceturtdienas stundu blokā")


# =========================================================== dokumenta veidnes
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def cell_text(cell, text, size=8.5, bold=False, color=None, align=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Calibri"
    if color:
        r.font.color.rgb = color


def para(doc, text, size=9, bold=False, color=None, before=0, after=3,
         italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = "Calibri"
    if color:
        r.font.color.rgb = color
    return p


def new_doc(klase, stundu_skaits, apraksts, grafiks=ADAZI):
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.2)
    s.top_margin = s.bottom_margin = Cm(1.0)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9)

    para(doc, "Fizika I  |  %s  |  %s  |  %s  |  3 stundas nedēļā  |  "
               "%d mācību stunda"
         % (grafiks.skola, klase, GADS, stundu_skaits),
         size=9, bold=True, color=NAVY, after=1)
    para(doc, apraksts, size=8.5, color=GREY, after=8)
    return doc


def pd1_piezime(p):
    """Teikums par pirmo pārbaudes darbu - datumu ņem no paša plāna, lai tas
    nekad nesarunātos ar kalendāru."""
    kods, _, _, datums, _ = p.vertejumi[0]
    diena = dt.datetime.strptime(datums, "%d.%m.%Y").date()
    ned = (diena - p.grafiks.sakums).days // 7 + 1
    return ("Pirmais pārbaudes darbs %s ir %s. - %d. mācību nedēļā."
            % (kods, datums, ned))


def kalendars(doc, vertejumi, piezime=None, grafiks=ADAZI):
    para(doc, "Darba organizācija un vērtēšanas kalendārs", size=13,
         bold=True, color=NAVY, after=3)
    para(doc, grafiks.apraksts(), size=8.5, color=GREY, after=6)
    if piezime:
        para(doc, piezime, size=8.5, italic=True, color=NAVY, after=6)

    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["Darbs", "Tēma / vērtēšanas objekts", "Svars",
                           "Datums", "Piezīme"]):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], h, size=9, bold=True, color=WHITE)
    for kods, tema, svars, datums, piez in vertejumi:
        row = t.add_row().cells
        cell_text(row[0], kods, bold=True, color=NAVY)
        cell_text(row[1], tema)
        cell_text(row[2], "%d %%" % svars, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[3], datums, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[4], piez, size=8, color=GREY)
    for c, w in enumerate([2.0, 9.5, 1.6, 2.8, 11.4]):
        for row in t.rows:
            row.cells[c].width = Cm(w)

    n_pd = sum(1 for v in vertejumi if v[0].startswith("PD"))
    n_ld = sum(1 for v in vertejumi if v[0].startswith("LD"))
    n_pr = sum(1 for v in vertejumi if v[0].startswith("PR"))
    para(doc, "Kopā %d vērtējumi: %d pārbaudes darbi (PD), %d laboratorijas "
              "darbi (LD) un %d prezentācija (PR). Svaru summa - 100 %%. "
              "PD vērtē tikai stundās faktiski apgūto saturu; neapgūtus "
              "izvēles padziļinājumus darbā neiekļauj."
         % (len(vertejumi), n_pd, n_ld, n_pr), size=8.5, before=6, after=10)


def temata_tabula(doc, virsraksts, komentars, rindas, padzilinajumi=None):
    para(doc, virsraksts, size=12, bold=True, color=NAVY, before=8, after=2)
    if komentars:
        para(doc, komentars, size=8.5, italic=True, color=GREY, after=4)

    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    hdr = ["N.p.k.", "Apakštemats", "Stundas tēma / jautājums",
           "Stundas sasniedzamais rezultāts", "Provizoriskais datums", "St."]
    for i, h in enumerate(hdr):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], h, size=8.5, bold=True, color=WHITE)
    for nr, apak, tema, sr, datums, st, veids in rindas:
        row = t.add_row().cells
        cell_text(row[0], nr, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[1], apak, bold=bool(veids))
        cell_text(row[2], tema, bold=bool(veids),
                  color=NAVY if veids else None)
        cell_text(row[3], sr)
        cell_text(row[4], datums, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[5], str(st), align=WD_ALIGN_PARAGRAPH.CENTER)
        if veids == "PD":
            for c in row:
                shade(c, "FFF4D6")
        elif veids in ("LD", "PR"):
            for c in row:
                shade(c, "E4F0DC")
    for c, w in enumerate([1.3, 4.0, 6.0, 11.0, 3.0, 1.0]):
        for row in t.rows:
            row.cells[c].width = Cm(w)

    if padzilinajumi:
        para(doc, "Izvēles padziļinājumi, ja atliek laiks: " + padzilinajumi,
             size=8.5, italic=True, color=GREY, before=3, after=6)


def noslegums(doc, virsraksts, komentars, rindas):
    para(doc, virsraksts, size=12, bold=True, color=NAVY, before=8, after=2)
    para(doc, komentars, size=8.5, italic=True, color=GREY, after=4)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["N.p.k.", "Saturs",
                           "Sasniedzamais rezultāts / izmantojums",
                           "Datums"]):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], h, size=8.5, bold=True, color=WHITE)
    for nr, apak, tema, sr, datums, st, veids in rindas:
        row = t.add_row().cells
        cell_text(row[0], nr, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[1], apak)
        cell_text(row[2], sr)
        cell_text(row[3], datums, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c, w in enumerate([1.3, 6.5, 15.5, 3.0]):
        for row in t.rows:
            row.cells[c].width = Cm(w)


# ============================================================ satura veidošana
#   ("st",  apakštemats, tēma, SR)                      -> 1 stunda
#   ("st2", apakštemats, tēma, SR)                      -> dubultstunda
#   ("PD",  kods, tēma, SR, svars, piezīme)             -> 1 stunda
#   ("LD",  kods, apakštemats, tēma, SR, svars, piez.)  -> dubultstunda
#   ("PR",  kods, apakštemats, tēma, SR, svars, piez.)  -> 1 stunda

DUBULTIE = ("st2", "LD")
VERTESANA = ("PD", "LD", "PR")        # summatīvie darbi

# Kāpēc stundu nedrīkst likt nākamajā slotā - katrai kaitei savs risinājums.
SASKELTS = "dubultstunda nesākas vienā dienā"
AIZNEMTA = "tajā dienā jau ir summatīvs darbs"
BLAKUS = "aiz tā tajā pašā dienā sanāktu vēl viens summatīvs darbs"


class Plans:
    """Piešķir stundu numurus un datumus; dubultstundas tur vienā dienā."""

    def __init__(self, grafiks=ADAZI):
        self.grafiks = grafiks
        self.slots = grafiks.slotu_saraksts()
        self.i = 0                    # nākamais brīvais slots
        self.n = 0                    # izsniegtais stundu numurs
        self.vertejumi = []
        self.mainas = []              # automātiskās pārkārtošanas
        self.aiznemtas = set()        # dienas, kurās jau ir summatīvs darbs

    # -- slotu izsniegšana --------------------------------------------------
    def _dubults_var(self):
        """Vai nākamās divas stundas ir vienā un tajā pašā dienā."""
        i = self.i
        return (i + 1 < len(self.slots)
                and self.slots[i][0] == self.slots[i + 1][0])

    def _nem_slotu(self, k):
        datums = self.slots[self.i][0]
        self.i += k
        self.n += k
        return datums

    def _garums(self, s):
        return 2 if s[0] in DUBULTIE else 1

    def _diena_aiznemta(self):
        """Vai nākamajā slotā jau paredzēts kāds summatīvs darbs."""
        return (self.i < len(self.slots)
                and self.slots[self.i][0] in self.aiznemtas)

    def _blakus_ta_pati_diena(self, s):
        """Vai uzreiz aiz šīs stundas tajā pašā dienā vēl paliek stunda."""
        j = self.i + self._garums(s)
        return (j < len(self.slots)
                and self.slots[j][0] == self.slots[self.i][0])

    def _kaite(self, rinda):
        """Kāpēc nākamo stundu nedrīkst likt nākamajā slotā (vai None).

        Dubultstunda nedrīkst sašķelties pa divām dienām, un vienā dienā
        nedrīkst sanākt divi summatīvi darbi - skolēnam nav jāraksta
        pārbaudes darbs uzreiz pēc laboratorijas darba. Tur, kur dienā ir
        trīs stundas pēc kārtas, otro sadursmi pamana tikai paskatoties
        vienu stundu uz priekšu.
        """
        s = rinda[0]
        if s[0] in DUBULTIE and not self._dubults_var():
            return SASKELTS
        if s[0] in VERTESANA and self._diena_aiznemta():
            return AIZNEMTA
        if (s[0] in VERTESANA and len(rinda) > 1
                and rinda[1][0] in VERTESANA
                and self._blakus_ta_pati_diena(s)):
            return BLAKUS
        return None

    def _pavelk(self, rinda):
        """Uz priekšu pavelk tuvāko parasto stundu; vai izdevās."""
        j = next((k for k, x in enumerate(rinda) if x[0] == "st"), None)
        if j is None:
            return False
        rinda.insert(0, rinda.pop(j))
        return True

    def _izlaist_dienu(self):
        """Dienas atlikušās stundas atstāj brīvas, lai darbs pārceļas uz
        nākamo dienu. Grafikā, kur dienā ir trīs stundas, stundu ir vairāk
        nekā plānā, un šis ir tas, kur rezerve noder."""
        diena = self.slots[self.i][0]
        while self.i < len(self.slots) and self.slots[self.i][0] == diena:
            self.i += 1
        return self.i < len(self.slots)

    def bloks(self, stundas):
        """Atgriež tabulas rindas; pati sakārto stundu secību tā, lai
        dubultstundas paliktu vienā dienā un summatīvie darbi nesakristu."""
        rinda = list(stundas)
        out = []
        while rinda:
            kaite = self._kaite(rinda)
            if kaite:
                nosaukums = rinda[0][1]
                if self._pavelk(rinda):
                    self.mainas.append(
                        "%d. stunda: «%s» - %s, uz priekšu pavilkta «%s»"
                        % (self.n + 1, nosaukums, kaite, rinda[0][1]))
                elif kaite == SASKELTS:
                    raise SystemExit("Nav ar ko aizpildīt dienu pirms "
                                     "dubultstundas")
                elif kaite == AIZNEMTA:
                    diena = d(self.slots[self.i][0])
                    self._izlaist_dienu()
                    self.mainas.append(
                        "«%s» pārcelts uz nākamo reizi - %s atlikusī stunda "
                        "paliek brīva" % (nosaukums, diena))
                else:
                    self.mainas.append(
                        "%d. stunda: «%s» - %s, un pārkārtot nav ar ko"
                        % (self.n + 1, nosaukums, kaite))
            s = rinda.pop(0)
            out.append(self._rinda(s))
        return out

    def _piez(self, teksts):
        """Piezīmē «{bloks}» aizstāj ar to, kur šajā grafikā ir dubultstunda -
        saturs par grafiku nezina, tāpēc to ieliek tikai izdrukājot."""
        return teksts.replace("{bloks}", self.grafiks.bloka_vieta)

    def _rinda(self, s):
        veids = s[0]
        if veids == "st":
            _, apak, tema, sr = s
            datums = self._nem_slotu(1)
            return ("%d." % self.n, apak, tema, sr, d(datums), 1, None)
        if veids == "st2":
            _, apak, tema, sr = s
            datums = self._nem_slotu(2)
            return ("%d.-%d." % (self.n - 1, self.n), apak, tema, sr,
                    d(datums), 2, None)
        if veids == "PD":
            _, kods, tema, sr, svars, piez = s
            datums = self._nem_slotu(1)
            self.vertejumi.append((kods, tema, svars, d(datums),
                                   self._piez(piez)))
            self.aiznemtas.add(datums)
            return ("%d." % self.n, "Summatīvā vērtēšana",
                    "%s: %s (%d %%)" % (kods, tema, svars), sr, d(datums), 1,
                    "PD")
        if veids == "LD":
            _, kods, apak, tema, sr, svars, piez = s
            datums = self._nem_slotu(2)
            self.vertejumi.append((kods, tema, svars, d(datums),
                                   self._piez(piez)))
            self.aiznemtas.add(datums)
            return ("%d.-%d." % (self.n - 1, self.n), apak,
                    "%s: %s (%d %%)" % (kods, tema, svars), sr, d(datums), 2,
                    "LD")
        if veids == "PR":
            _, kods, apak, tema, sr, svars, piez = s
            datums = self._nem_slotu(1)
            self.vertejumi.append((kods, tema, svars, d(datums),
                                   self._piez(piez)))
            self.aiznemtas.add(datums)
            return ("%d." % self.n, apak,
                    "%s: %s (%d %%)" % (kods, tema, svars), sr, d(datums), 1,
                    "PR")
        raise SystemExit("Nezināms stundas veids: %r" % (veids,))

    # -- pārbaudes ----------------------------------------------------------
    def parbaudi(self):
        dienas = [v[3] for v in self.vertejumi]
        divreiz = sorted({x for x in dienas if dienas.count(x) > 1})
        assert not divreiz, ("Vienā dienā sanāk divi summatīvi darbi: %s"
                             % ", ".join(divreiz))
        svars = sum(v[2] for v in self.vertejumi)
        assert svars == 100, ("Vērtējumu svaru summa ir %d %%, jābūt 100 %%"
                              % svars)
        assert self.i <= len(self.slots), (
            "Plānā ir %d stundas, bet kalendārā tikai %d"
            % (self.i, len(self.slots)))
