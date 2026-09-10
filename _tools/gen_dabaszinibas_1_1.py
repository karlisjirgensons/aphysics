# -*- coding: utf-8 -*-
"""
Ģenerē prezentāciju: Dabaszinības (fizikas daļa), 10. klase
Temats 1. "Pasaule ap mums un tās pētīšana"
Stunda 1.1. "Matērija - viela un lauks. Mikropasaule, makropasaule, megapasaule"

Noformējums atbilst rules_dabaszinibas.txt prasībām:
  - vispirms īsa teorija, tad uzdevumi;
  - uzdevumu slaidos vienlaikus redzami 3 uzdevumi (1. augšā, zem atstarpes vēl 2);
  - atstarpē soli pa solim aug risinājums, uzdevumi paliek redzami;
  - risinājums latviešu standartā: Dots / Jāaprēķina / Formulas / Aprēķins / Atbilde.

Risinājuma atklāšanas princips:
  - katrs JAUNAIS ieraksts parādās liels un centrēts uzmanības kartītē;
  - nākamajā slaidā tas sarūk un nostājas savā vietā (Dots / Jāaprēķina /
    Formulas joslā augšā, aprēķina soļi - sarakstā zem tās);
  - ja aprēķina soļiem pietrūkst vietas, vecākie soļi tiek noņemti (paliek "...").

Viss teksts tiek automātiski samazināts tā, lai neizietu ārpus sava rāmja.
"""

import math
import re
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import mathfmt as MF
import mathfmt_prose as MP

# ------------------------------------------------------------------ palete
FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
LIGHTBLUE = RGBColor(0xEA, 0xF2, 0xFB)
GOLD = RGBColor(0xB3, 0x86, 0x00)
LIGHTGOLD = RGBColor(0xFF, 0xF6, 0xDF)
GREY = RGBColor(0x60, 0x60, 0x60)
LINEGREY = RGBColor(0xD0, 0xD0, 0xD0)
LIGHTGREY = RGBColor(0xF5, 0xF5, 0xF5)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
RED = RGBColor(0xB4, 0x1E, 0x1E)
LIGHTRED = RGBColor(0xFD, 0xEE, 0xEE)
GREEN = RGBColor(0x1E, 0x6B, 0x3A)
LIGHTGREEN = RGBColor(0xE9, 0xF4, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = 13.333, 7.5
MX = 0.55                      # lapas kreisā/labā mala
CW = SW - 2 * MX               # satura platums = 12.233"

# teksta augstuma novērtējums (konservatīvs, lai teksts nekad neizietu ārā)
CHAR_W = 0.50
LINE_H = 1.24


def est_h(lines, width_in):
    """Novērtē teksta bloka augstumu collās dotajā platumā."""
    w_pt = width_in * 72.0
    total = 0.0
    for ln in lines:
        sz = float(ln.get("size", 20))
        cpl = max(1, int(w_pt / (sz * CHAR_W)))
        n = max(1, math.ceil(len(ln["t"]) / cpl))
        total += n * sz * LINE_H + ln.get("space", 0) + ln.get("after", 0)
    return total / 72.0


def fit(lines, width_in, height_in, min_size=8.0):
    """Samazina fontus, līdz bloks ietilpst dotajā laukumā."""
    k, scaled = 1.0, None
    while k > 0.30:
        scaled = [dict(l, size=max(min_size, float(l.get("size", 20)) * k),
                       space=l.get("space", 0) * k) for l in lines]
        if est_h(scaled, width_in) <= height_in:
            return scaled
        k -= 0.03
    return scaled


# ------------------------------------------------------------------ pamati
def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def box(slide, x, y, w, h, fill=None, line=None, lw=1.25, rounded=True,
        name=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.05
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def rule(slide, x, y, w, color=LINEGREY, h=0.015, name="RULE"):
    box(slide, x, y, w, h, fill=color, line=None, rounded=False, name=name)


# ------------------------------------------------------- bultas un nogriežņi
# Vektoru zīmējumi. Ģeometriju glabā formas nosaukumā (ARR:x1,y1,x2,y2,...),
# tāpēc HTML zīmētājs to var atkārtot precīzi tādu pašu (DRY) - collās, tāpat
# kā slaidā.

ARROW_LW = 0.030          # bultas kāta biezums collās
ARROW_HEAD = 0.20         # uzgaļa garums collās


def _arrow_pts(x1, y1, x2, y2, lw, head):
    """Bultas kontūra: kāts + trīsstūra uzgalis (vai tikai nogrieznis)."""
    dx, dy = x2 - x1, y2 - y1
    ln = (dx * dx + dy * dy) ** 0.5 or 1e-6
    ux, uy = dx / ln, dy / ln
    px, py = -uy, ux                       # perpendikuls
    hl = min(head, ln * 0.34) if head else 0.0
    hw = lw * 2.2
    sw = lw / 2.0

    def pt(u, v):
        return (x1 + ux * u + px * v, y1 + uy * u + py * v)

    if not head:
        return [pt(0, -sw), pt(ln, -sw), pt(ln, sw), pt(0, sw)]
    return [pt(0, -sw), pt(ln - hl, -sw), pt(ln - hl, -hw), pt(ln, 0),
            pt(ln - hl, hw), pt(ln - hl, sw), pt(0, sw)]


def arrow(slide, x1, y1, x2, y2, color=BLUE, lw=ARROW_LW, head=True,
          name=None):
    """Vektora bulta no (x1, y1) uz (x2, y2) collās."""
    pts = _arrow_pts(x1, y1, x2, y2, lw, head)
    bld = slide.shapes.build_freeform(pts[0][0], pts[0][1], Inches(1).emu)
    bld.add_line_segments(pts[1:], close=True)
    shp = bld.convert_to_shape()
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    shp.name = "%s|ARR:%.3f,%.3f,%.3f,%.3f,%.3f,%d" % (
        name or "FIG", x1, y1, x2, y2, lw, 1 if head else 0)
    return shp


def segment(slide, x1, y1, x2, y2, color=LINEGREY, lw=0.014, name=None):
    """Palīglīnija bez uzgaļa (paralelograma malas, projekcijas)."""
    shp = arrow(slide, x1, y1, x2, y2, color, lw, head=False, name=name)
    shp.name = shp.name.replace("|ARR:", "|SEG:", 1)
    return shp


def txt(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


# Vektora bultiņas (U+20D7) Calibri fontā nav, tāpēc PowerPoint to aizstāj
# ar citu fontu un bultiņa nostājas blakus burtam, nevis virs tā. Simbolu
# fontā burts un bultiņa ir kopā, tāpēc vektora simbolu raksta ar to.
VEC_FONT = "Segoe UI Symbol"


def _write_runs(p, text, size, bold, italic, color):
    """Teksta gabalus ieraksta rindkopā; vektoram - vektora fontu.

    Viss teksts .pptx failos iet caur šo funkciju (DRY), tāpēc bultiņas
    labojums der visām prezentācijām uzreiz.
    """
    for kind, chunk in MF.split_vectors(text) or [("t", "")]:
        r = p.add_run()
        r.text = chunk + (MF.VEC_MARK if kind == "v" else "")
        f = r.font
        f.name = VEC_FONT if kind == "v" else FONT
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.color.rgb = color


def put(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, autofit=True,
        name=None):
    """Ieraksta tekstu, vajadzības gadījumā samazinot to, lai ietilptu."""
    if autofit:
        lines = fit(lines, w, h)
    tf = txt(slide, x, y, w, h, anchor, name=name)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(ln.get("space", 0))
        p.space_after = Pt(ln.get("after", 0))
        _write_runs(p, ln["t"], ln.get("size", 20), ln.get("bold", False),
                    ln.get("italic", False), ln.get("color", DARK))
    return tf


_UID = [0]


def uid(prefix):
    _UID[0] += 1
    return "%s%d" % (prefix, _UID[0])


# Cik daudz parasta teksta drīkst būt rindā, lai to vēl zīmētu formulu
# dzinējs: formulu rindu ("f = 1/T") liek vertikāli, veselu teikumu ar daļu
# vidū - ne, jo tas jāaplauž pa vārdiem.
MATH_LINE_TEXT = 30
MATH_LINE_MAX = 28                      # īsa rinda bez vārdiem = tīra formula
_WORD = re.compile(r"[A-Za-zĀ-žĀ-ū]{4,}")


def math_atoms(text):
    """Rindas atomi ar to pašu izšķiršanu, ko lieto math_line().

    Īsu formulu rindu ("T = t/N.") drīkst lasīt agresīvi - tur nav vārdu,
    ko sajaukt ar mērvienībām. Teikumam lieto piesardzīgo parsētāju.
    """
    s = text.strip()
    if len(s) <= MATH_LINE_MAX and not _WORD.search(s):
        return MF.parse_math(s)
    return MP.parse_prose(s)


def math_line(text):
    """Vai rinda ir formula (nevis teikums), kas jāzīmē vertikāli."""
    atoms = math_atoms(text)
    if not any(a[0] in ("f", "r") for a in atoms):
        return False
    return sum(len(a[1]) for a in atoms if a[0] == "t") <= MATH_LINE_TEXT


def _scaled(lines, k):
    """Rindu kopija ar samazinātu fontu un atstarpēm."""
    return [dict(ln, size=float(ln.get("size", 20)) * k,
                 space=ln.get("space", 0) * k,
                 math=math_line(ln["t"])) for ln in lines]


def panel(slide, x, y, w, h, lines, accent=BLUE, fill=WHITE, pad=0.24,
          anchor=MSO_ANCHOR.TOP, lw=1.25, name=None):
    """Rāmis ar tekstu, kas vienmēr ietilpst rāmī.

    rules_lessons.txt prasa vertikālas daļas arī teorijā, tāpēc formulu
    rindas zīmē formulu dzinējs; teikumus - parastais teksta ceļš ar
    automātisku samazināšanu.
    """
    n = name or uid("PANEL")
    box(slide, x, y, w, h, fill=fill, line=accent, lw=lw, name="BOX:" + n)
    ix, iy = x + pad, y + pad * 0.62
    iw, ih = w - 2 * pad, h - 1.24 * pad
    if any(math_line(ln["t"]) for ln in lines):
        k = 1.0
        while k > 0.45 and _lines_h(_scaled(lines, k), iw) > ih:
            k -= 0.04
        _put_lines(slide, ix, iy, iw, _scaled(lines, k), name="TXT:" + n)
        return
    put(slide, ix, iy, iw, ih, lines, anchor=anchor, name="TXT:" + n)


# ============================================ vertikālas daļas (a/b -> a virs b)
# rules_lessons.txt: dalījumu raksta ar skaitītāju virs saucēja, nevis rindā.
# Mērvienības (m/s, kg/m³, N/m²) paliek rindā - tā ir latviešu standarta forma.

# Formulas izmēri (fonta izmēra daļās) dzīvo mathfmt - tos pašus lieto arī
# HTML zīmētājs, tāpēc slaids un lapa izskatās vienādi (DRY).
from mathfmt import (                          # noqa: E402
    MATH_LH, MATH_PAD, MATH_H, MATH_ROOT_LH, ROOT_SIDE, ROOT_ASC, ROOT_DESC,
    ROOT_BAR, ROOT_LEAD, ROOT_INNER, BASE_OFF, ROOT_TOP, root_extent)


def root_sign_w(atoms, size):
    """Saknes zīmes platums punktos - tas seko zīmes augstumam."""
    up, dn = root_extent(atoms)
    return MF.root_width(up + dn) * size


def _atom_w(a, size, bold=False, italic=False):
    """Viena atoma platums punktos."""
    if a[0] == "t":
        return MF.text_w(a[1], size, bold, italic)
    if a[0] == "r":
        return ((ROOT_LEAD + ROOT_INNER + ROOT_SIDE) * size
                + root_sign_w(a[1], size)
                + sum(_atom_w(x, size, bold, italic) for x in a[1]))
    return (max(MF.text_w(a[1], size, bold, italic),
                MF.text_w(a[2], size, bold, italic))
            + 0.60 * size)               # elpa abās pusēs daļai


def math_w(text, size, bold=False, italic=False, atoms=None):
    """Rindas platums collās, ja to zīmē ar vertikālām daļām un saknēm."""
    if atoms is None:
        atoms = MF.parse_math(text)
    return sum(_atom_w(a, size, bold, italic) for a in atoms) / 72.0


def math_h(text, size, atoms=None):
    """Rindas augstums collās (ar daļu tas ir gandrīz trīskāršs).

    Augstumu mēra simetriski ap rindas viduslīniju - tur put_math() liek
    daļas svītru -, tāpēc te ir divkāršs lielākais izlēciens uz augšu vai
    leju.
    """
    if atoms is None:
        atoms = MF.parse_math(text)
    if any(MF.has_root_fraction(a) for a in atoms):
        return (MATH_H + 2 * ROOT_TOP) * size / 72.0
    if any(a[0] == "f" for a in atoms):
        return MATH_H * size / 72.0
    if any(a[0] == "r" for a in atoms):
        return MATH_ROOT_LH * size / 72.0
    return MATH_LH * size / 72.0


def math_fits(text, w_in, size, bold=False):
    return math_w(text, size, bold) <= w_in


def _mline(slide, x, y, w, text, size, color, bold, italic, name=None,
           align=PP_ALIGN.CENTER):
    """Viena teksta gabala zīmēšana bez aplaušanas."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w),
                                  Inches(MATH_LH * size / 72.0))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    _write_runs(p, text, size, bold, italic, color)
    return tb


def root_sign(slide, x, y, h_em, size, color, name=None):
    """Saknes zīme kā daudzstūris (MF.root_pts) - to būvē pēc augstuma."""
    k = Inches(size / 72.0).emu
    _, pts = MF.root_pts(h_em)
    bld = slide.shapes.build_freeform(pts[0][0], pts[0][1], k)
    bld.add_line_segments(pts[1:], close=True)
    shp = bld.convert_to_shape(Inches(x), Inches(y))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def _draw_atoms(slide, atoms, cx, ym, ctx, sub=""):
    """Uzzīmē atomu virkni pa kreisi no cx; atgriež jauno cx.

    Viena vieta, kur zināms, kā izskatās teksts, daļa un sakne (SRP) -
    sakne izsauc to pašu funkciju saviem apakšatomiem, tāpēc daļa zem
    vinkula sanāk pati no sevis (DRY).
    """
    size, color = ctx["size"], ctx["color"]
    bold, italic = ctx["bold"], ctx["italic"]
    lh, pad = ctx["lh"], ctx["pad"]

    def nm(role):
        return "%sMATH:%s:%s%s" % (ctx["pfx"], ctx["mid"], sub, role)

    for a in atoms:
        if a[0] == "t":
            tw = MF.text_w(a[1], size, bold, italic) / 72.0
            if tw <= 0:
                continue
            _mline(slide, cx, ym - lh / 2.0, tw + 0.02, a[1], size, color,
                   bold, italic, name=nm("t"),
                   align=PP_ALIGN.LEFT if sub else PP_ALIGN.CENTER)
            cx += tw
        elif a[0] == "r":
            # saknes zīme aptver visu izteiksmi; vinkuls turpina zīmes
            # augšmalu, tāpēc tas sākas tieši tur, kur zīme beidzas
            up, dn = root_extent(a[1])
            top = ym - up * size / 72.0
            rw = root_sign_w(a[1], size) / 72.0
            inner = ROOT_INNER * size / 72.0
            ew = sum(_atom_w(x, size, bold, italic) for x in a[1]) / 72.0
            cx += ROOT_LEAD * size / 72.0
            tall = any(z[0] == "f" for z in a[1])
            root_sign(slide, cx, top, up + dn, size, color,
                      name=nm("rgT" if tall else "rg"))
            _draw_atoms(slide, a[1], cx + rw + inner, ym, ctx, sub="r")
            box(slide, cx + rw * 0.94, top,
                rw * 0.06 + inner + ew + 0.02, ROOT_BAR * size / 72.0,
                fill=color, line=None, rounded=False, name=nm("rv"))
            cx += rw + inner + ew + ROOT_SIDE * size / 72.0
        else:
            num, den = a[1], a[2]
            fw = max(MF.text_w(num, size, bold, italic),
                     MF.text_w(den, size, bold, italic)) / 72.0
            side = 0.30 * size / 72.0
            cx += side
            _mline(slide, cx, ym - pad / 2.0 - lh, fw, num, size, color,
                   bold, italic, name=nm("n"))
            box(slide, cx, ym - 0.008, fw, 0.016, fill=color, line=None,
                rounded=False, name=nm("b"))
            _mline(slide, cx, ym + pad / 2.0, fw, den, size, color, bold,
                   italic, name=nm("d"))
            cx += fw + side
    return cx


def put_math(slide, x, y, w, h, text, size, color=DARK, bold=False,
             italic=False, align=PP_ALIGN.CENTER, min_size=9.0, name=None,
             atoms=None):
    """Uzzīmē vienu formulas rindu, dalījumus liekot vertikāli.

    Ja tekstā daļu nav, uzvedība ir tāda pati kā ar put() - viena rinda.
    Atgriež faktisko fonta izmēru.
    """
    if atoms is None:
        atoms = MF.parse_math(text)
    if not any(a[0] in ("f", "r") for a in atoms):
        put(slide, x, y, w, h, [{"t": text, "size": size, "bold": bold,
                                 "italic": italic, "color": color,
                                 "align": align}],
            anchor=MSO_ANCHOR.MIDDLE, name=name)
        return size
    mid = uid("M")
    pfx = (name + "|") if name else ""

    # samazina, līdz ietilpst platumā un augstumā
    # Augstumu prasa tik, cik rindai tiešām vajag: saknes rindai pietiek ar
    # vienu rindu un vinkulu, daļai - ar divām. Agrāk abām prasīja daļas
    # augstumu, tāpēc saknes formulas sarāvās līdz min_size.
    while size > min_size:
        if (math_w(text, size, bold, italic, atoms) <= w
                and math_h(text, size, atoms=atoms) <= h):
            break
        size -= 0.5
    size = max(size, min_size)

    lh = MATH_LH * size / 72.0
    pad = MATH_PAD * size / 72.0
    total = math_w(text, size, bold, italic, atoms)
    if align == PP_ALIGN.CENTER:
        cx = x + (w - total) / 2.0
    elif align == PP_ALIGN.RIGHT:
        cx = x + w - total
    else:
        cx = x
    cx = max(x, cx)
    ym = y + h / 2.0                       # daļas svītras līmenis

    ctx = dict(size=size, color=color, bold=bold, italic=italic,
               pfx=pfx, mid=mid, lh=lh, pad=pad)
    _draw_atoms(slide, atoms, cx, ym, ctx)
    return size


KICKER = "DABASZINĪBAS · 10. KLASE · 1. TEMATS: PASAULE AP MUMS UN TĀS PĒTĪŠANA"
FOOT = "1.1. Matērija. Viela un lauks. Mikropasaule, makropasaule, megapasaule"


def header(slide, title, right=None):
    put(slide, MX, 0.24, 9.4, 0.24,
        [{"t": KICKER, "size": 10.5, "bold": True, "color": GREY}],
        autofit=False, name="KICKER")
    put(slide, MX, 0.48, 9.4, 0.44,
        [{"t": title, "size": 24, "bold": True, "color": NAVY}],
        name="TITLE")
    if right:
        put(slide, 10.1, 0.52, CW - 9.55, 0.36,
            [{"t": right, "size": 14, "bold": True, "color": GOLD,
              "align": PP_ALIGN.RIGHT}], autofit=False, name="BADGE")
    rule(slide, MX, 1.00, CW, GOLD, 0.022)


def footer(slide, text=None):
    # FOOT nomaina katrai stundai, tāpēc to nolasa izsaukuma brīdī,
    # nevis funkcijas definēšanas brīdī.
    put(slide, MX, SH - 0.40, CW, 0.24,
        [{"t": FOOT if text is None else text, "size": 10.5, "color": GREY}],
        autofit=False, name="FOOTER")


# ============================================================ TEORIJAS SLAIDI
def slide_title(prs):
    s = blank(prs)
    box(s, 0, 0, SW, SH, fill=NAVY, rounded=False)
    box(s, 0, 6.10, SW, 1.40, fill=RGBColor(0x17, 0x2B, 0x4D), rounded=False)

    put(s, 1.05, 1.45, CW, 0.32,
        [{"t": "DABASZINĪBAS · 10. KLASE · FIZIKAS DAĻA", "size": 16,
          "bold": True, "color": RGBColor(0x9D, 0xC3, 0xE6)}], autofit=False)
    put(s, 1.05, 2.00, 11.0, 0.70,
        [{"t": "1. temats. Pasaule ap mums un tās pētīšana", "size": 36,
          "bold": True, "color": WHITE}])
    rule(s, 1.05, 3.05, 3.60, GOLD, 0.03)

    put(s, 1.05, 3.45, 11.0, 2.10, [
        {"t": "1.1. stunda", "size": 21, "bold": True, "color": GOLD},
        {"t": "Matērija — viela un lauks", "size": 31, "bold": True,
         "color": WHITE, "space": 12},
        {"t": "Mikropasaule · Makropasaule · Megapasaule", "size": 31,
         "bold": True, "color": WHITE, "space": 6},
    ])
    put(s, 1.05, 6.50, 11.0, 0.55,
        [{"t": "Fizikālie lielumi · SI mērvienības · Priedēkļi · "
               "Uzdevuma risinājuma pieraksts", "size": 15,
          "color": RGBColor(0xBD, 0xD7, 0xEE)}])
    return s


def slide_merkis(prs):
    s = blank(prs)
    header(s, "Stundas mērķis un sasniedzamais rezultāts")

    panel(s, MX, 1.20, CW, 1.15, [
        {"t": "MĒRĶIS", "size": 12, "bold": True, "color": GOLD},
        {"t": "Izprast, kas ir matērija, un salīdzināt mikropasaules, "
              "makropasaules un megapasaules objektus, lietojot fizikālos "
              "lielumus un SI mērvienības.", "size": 20, "space": 6},
    ], accent=GOLD, fill=LIGHTGOLD)

    put(s, MX, 2.62, CW, 0.32,
        [{"t": "Stundas beigās es protu:", "size": 16, "bold": True,
          "color": NAVY}], autofit=False)

    items = [
        "skaidrot matēriju, lietojot jēdzienus “viela” un “lauks”;",
        "klasificēt objektus mikro-, makro- un megapasaulē pēc to izmēra;",
        "nosaukt katras pasaules pētīšanas metodes un mērierīces;",
        "lietot fizikālo lielumu apzīmējumus, SI mērvienības un priedēkļus;",
        "pierakstīt risinājumu: Dots → Jāaprēķina → Formulas → Aprēķins "
        "→ Atbilde.",
    ]
    y = 3.06
    for i, it in enumerate(items, 1):
        box(s, MX, y, 0.44, 0.44, fill=BLUE, line=None)
        put(s, MX + 0.10, y + 0.09, 0.30, 0.28,
            [{"t": str(i), "size": 15, "bold": True, "color": WHITE,
              "align": PP_ALIGN.CENTER}], autofit=False)
        put(s, MX + 0.68, y + 0.05, CW - 0.80, 0.42,
            [{"t": it, "size": 19}])
        y += 0.66
    footer(s)
    return s


def slide_materija(prs):
    s = blank(prs)
    header(s, "Matērija un tās eksistences formas")

    panel(s, MX, 1.20, CW, 0.86, [
        {"t": "MATĒRIJA — viss, kas pastāv objektīvi, neatkarīgi no cilvēka "
              "apziņas. Tā pastāv divās eksistences formās:", "size": 19,
         "bold": True, "color": NAVY},
    ], accent=LINEGREY, fill=LIGHTGREY, anchor=MSO_ANCHOR.MIDDLE)

    colw = (CW - 0.38) / 2
    panel(s, MX, 2.22, colw, 3.32, [
        {"t": "VIELA", "size": 26, "bold": True, "color": BLUE},
        {"t": "Sastāv no daļiņām — atomiem, molekulām, joniem.",
         "size": 17, "space": 10},
        {"t": "Raksturlielums:  masa  m  [kg]", "size": 19, "bold": True,
         "space": 12},
        {"t": "Var pataustīt, nosvērt, izmērīt tilpumu.", "size": 15,
         "space": 8, "color": GREY},
        {"t": "•  ūdens, dzelzs, gaiss, stikls", "size": 17, "space": 14},
        {"t": "•  cilvēka ķermenis, Zeme, zvaigznes", "size": 17, "space": 4},
        {"t": "ρ = m / V   — vielas blīvums [kg/m³]", "size": 18,
         "bold": True, "space": 14, "color": NAVY},
    ], accent=BLUE)

    panel(s, MX + colw + 0.38, 2.22, colw, 3.32, [
        {"t": "LAUKS", "size": 26, "bold": True, "color": GOLD},
        {"t": "Nesastāv no daļiņām; nodrošina mijiedarbību attālumā.",
         "size": 17, "space": 10},
        {"t": "Raksturlielums:  lauka stiprums", "size": 19, "bold": True,
         "space": 12},
        {"t": "Nevar pataustīt — nosaka pēc iedarbības.", "size": 15,
         "space": 8, "color": GREY},
        {"t": "•  gravitācijas lauks (Zeme pievelk)", "size": 17,
         "space": 14},
        {"t": "•  elektriskais un magnētiskais lauks", "size": 17,
         "space": 4},
        {"t": "F = mg   — smaguma spēks gravitācijas laukā", "size": 18,
         "bold": True, "space": 14, "color": NAVY},
    ], accent=GOLD)

    panel(s, MX, 5.70, CW, 1.02, [
        {"t": "KOPĪGAIS:  abām matērijas formām piemīt enerģija, un noteiktos "
              "apstākļos tās var pārvērsties viena otrā", "size": 18,
         "bold": True, "color": NAVY},
        {"t": "(piemēram, elektronam un pozitronam anihilējot, viela "
              "pārvēršas elektromagnētiskajā starojumā).", "size": 15,
         "color": GREY, "space": 5},
    ], accent=NAVY)
    footer(s)
    return s


def slide_pasaules(prs):
    s = blank(prs)
    header(s, "Pasaules nosacītais iedalījums pēc objektu izmēra")

    cols = [
        ("MIKROPASAULE", "mazāk par 10⁻⁷ m",
         ["elementārdaļiņas  ~10⁻¹⁵ m", "atoma kodols  ~10⁻¹⁴ m",
          "atoms  ~10⁻¹⁰ m", "molekula  ~10⁻⁹ m", "vīruss, baktērija"],
         ["optiskais mikroskops", "elektronmikroskops", "atomspēku mikroskops"],
         BLUE),
        ("MAKROPASAULE", "no 10⁻⁷ m līdz ~10⁷ m",
         ["smilšu grauds  ~10⁻⁴ m", "cilvēks  ~1,7 m", "māja, automašīna",
          "kalni, jūras", "Zeme  R = 6,37·10⁶ m"],
         ["lineāls, bīdmērs, svari", "termometrs, hronometrs",
          "digitālie sensori"],
         GOLD),
        ("MEGAPASAULE", "vairāk par ~10⁷ m",
         ["planētas un pavadoņi", "Saule  R = 6,96·10⁸ m",
          "zvaigžņu sistēmas", "galaktikas  ~10²¹ m", "Visums  ~10²⁶ m"],
         ["optiskie teleskopi", "radioteleskopi", "kosmiskās zondes"],
         GREEN),
    ]

    colw = (CW - 2 * 0.30) / 3
    x = MX
    for name, rng, objs, tools, c in cols:
        lines = [
            {"t": name, "size": 21, "bold": True, "color": c},
            {"t": rng, "size": 16, "bold": True, "space": 4},
            {"t": "OBJEKTI", "size": 11, "bold": True, "space": 14,
             "color": GREY},
        ]
        for o in objs:
            lines.append({"t": "•  " + o, "size": 16, "space": 4})
        lines.append({"t": "PĒTA AR", "size": 11, "bold": True, "space": 14,
                      "color": GREY})
        for t in tools:
            lines.append({"t": "•  " + t, "size": 16, "space": 4})
        panel(s, x, 1.20, colw, 5.06, lines, accent=c)
        x += colw + 0.30

    panel(s, MX, 6.40, CW, 0.66, [
        {"t": "Iedalījums ir NOSACĪTS — robežas nav stingri noteiktas. "
              "Katras pasaules pētīšanai vajadzīgas savas metodes un "
              "mērierīces.", "size": 17, "bold": True, "color": NAVY},
    ], accent=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    return s


def slide_limeni(prs):
    s = blank(prs)
    header(s, "Pasaules organizācijas līmeņi")

    put(s, MX, 1.18, CW, 0.30,
        [{"t": "No mazākā uz lielāko — katrs nākamais līmenis veidojas no "
               "iepriekšējā:", "size": 17, "color": GREY}], autofit=False)

    steps = [
        ("atoms", "10⁻¹⁰ m", BLUE), ("molekula", "10⁻⁹ m", BLUE),
        ("šūna", "10⁻⁵ m", BLUE), ("audi", "10⁻³ m", GOLD),
        ("orgāns", "10⁻¹ m", GOLD), ("organisms", "10⁰ m", GOLD),
        ("populācija", "10² m", GOLD), ("ekosistēma", "10⁴ m", GOLD),
        ("biosfēra", "10⁷ m", GREEN), ("planēta", "10⁷ m", GREEN),
        ("zvaigžņu sist.", "10¹³ m", GREEN), ("Galaktika", "10²¹ m", GREEN),
        ("Visums", "10²⁶ m", GREEN),
    ]
    per_row, gapx = 5, 0.26
    w = (CW - (per_row - 1) * gapx) / per_row
    h = 0.92
    x, y = MX, 1.62
    for i, (n, sc, c) in enumerate(steps):
        if i and i % per_row == 0:
            x = MX
            y += h + 0.32
        box(s, x, y, w, h, fill=WHITE, line=c, lw=1.5)
        put(s, x + 0.10, y + 0.14, w - 0.20, h - 0.28, [
            {"t": n, "size": 15, "bold": True, "color": c,
             "align": PP_ALIGN.CENTER},
            {"t": "≈ " + sc, "size": 14, "color": GREY, "space": 2,
             "align": PP_ALIGN.CENTER},
        ])
        if (i + 1) % per_row != 0 and i != len(steps) - 1:
            put(s, x + w, y + 0.26, gapx, 0.40,
                [{"t": "›", "size": 20, "bold": True, "color": LINEGREY,
                  "align": PP_ALIGN.CENTER}], autofit=False)
        x += w + gapx

    panel(s, MX, 5.66, CW, 0.80, [
        {"t": "Visi līmeņi ir savstarpēji saistīti: izmaiņas vienā līmenī "
              "ietekmē pārējos — piemēram, molekulu bojājumi šūnā ietekmē "
              "visu organismu.", "size": 17, "bold": True, "color": NAVY},
    ], accent=LINEGREY, fill=LIGHTGREY, anchor=MSO_ANCHOR.MIDDLE)
    footer(s)
    return s


def slide_si(prs):
    s = blank(prs)
    header(s, "Fizikālie lielumi un SI pamatvienības")

    rows = [
        ("garums", "l, s, h, d", "metrs", "m"),
        ("masa", "m", "kilograms", "kg"),
        ("laiks", "t", "sekunde", "s"),
        ("elektriskās strāvas stiprums", "I", "ampērs", "A"),
        ("temperatūra", "T", "kelvins", "K"),
        ("vielas daudzums", "n", "mols", "mol"),
        ("gaismas stiprums", "Iᵥ", "kandela", "cd"),
    ]
    hdr = ["Fizikālais lielums", "Apzīmējums", "SI vienība", "Simbols"]
    ws = [4.80, 2.30, 2.80, 2.333]
    xs, acc = [], MX
    for w in ws:
        xs.append(acc)
        acc += w

    for hh, xx, ww in zip(hdr, xs, ws):
        box(s, xx, 1.20, ww, 0.52, fill=NAVY, line=None, rounded=False)
        put(s, xx + 0.14, 1.31, ww - 0.28, 0.32,
            [{"t": hh, "size": 15, "bold": True, "color": WHITE}],
            autofit=False)

    y = 1.76
    for r_i, row in enumerate(rows):
        bg = WHITE if r_i % 2 == 0 else LIGHTBLUE
        for c_i, (val, xx, ww) in enumerate(zip(row, xs, ws)):
            box(s, xx, y, ww, 0.55, fill=bg, line=LINEGREY, lw=0.75,
                rounded=False)
            put(s, xx + 0.14, y + 0.12, ww - 0.28, 0.36,
                [{"t": val, "size": 18, "bold": (c_i in (1, 3)),
                  "color": NAVY if c_i in (1, 3) else DARK}])
        y += 0.55

    panel(s, MX, 5.90, CW, 1.02, [
        {"t": "ATCERIES!  Pirms aprēķina visi lielumi jāizsaka SI vienībās.",
         "size": 19, "bold": True, "color": GOLD},
        {"t": "540 g = 0,54 kg   ·   200 cm³ = 2·10⁻⁴ m³   ·   "
              "72 km/h = 20 m/s   ·   25 °C = 298 K", "size": 18, "space": 6},
    ], accent=GOLD, fill=LIGHTGOLD)
    footer(s)
    return s


def slide_priedekli(prs):
    s = blank(prs)
    header(s, "Priedēkļi mērvienību daudzkārtņu un daļvienību veidošanai")

    big = [("tera", "T", "10¹²"), ("giga", "G", "10⁹"),
           ("mega", "M", "10⁶"), ("kilo", "k", "10³"),
           ("hekto", "h", "10²"), ("deka", "da", "10¹")]
    small = [("deci", "d", "10⁻¹"), ("centi", "c", "10⁻²"),
             ("mili", "m", "10⁻³"), ("mikro", "µ", "10⁻⁶"),
             ("nano", "n", "10⁻⁹"), ("piko", "p", "10⁻¹²")]

    colw = (CW - 0.38) / 2

    def col(x, title, data, c, bg):
        box(s, x, 1.20, colw, 4.36, fill=bg, line=c, lw=1.25)
        put(s, x + 0.28, 1.34, colw - 0.56, 0.34,
            [{"t": title, "size": 17, "bold": True, "color": c}],
            autofit=False)
        y = 1.82
        for name, sym, val in data:
            put(s, x + 0.28, y, 2.10, 0.38, [{"t": name, "size": 19}],
                autofit=False)
            put(s, x + 2.50, y, 0.90, 0.38,
                [{"t": sym, "size": 19, "bold": True, "color": c}],
                autofit=False)
            put(s, x + 3.60, y, colw - 3.90, 0.38,
                [{"t": val, "size": 19, "bold": True}], autofit=False)
            y += 0.60

    col(MX, "DAUDZKĀRTNES — lielākas vienības", big, BLUE, LIGHTBLUE)
    col(MX + colw + 0.38, "DAĻVIENĪBAS — mazākas vienības", small, GOLD,
        LIGHTGOLD)

    panel(s, MX, 5.74, CW, 1.20, [
        {"t": "PIEMĒRI", "size": 12, "bold": True, "color": NAVY},
        {"t": "1 nm = 10⁻⁹ m   ·   1 µm = 10⁻⁶ m   ·   1 mm = 10⁻³ m   ·   "
              "1 km = 10³ m", "size": 19, "space": 7},
        {"t": "0,28 nm = 0,28 · 10⁻⁹ m = 2,8 · 10⁻¹⁰ m", "size": 19,
         "bold": True, "space": 7, "color": RED},
    ], accent=NAVY)
    return s


def slide_pieraksts(prs):
    s = blank(prs)
    header(s, "Kā pieraksta uzdevuma risinājumu")

    put(s, MX, 1.18, CW, 0.30,
        [{"t": "Paraugs: metāla detaļas masa 540 g, tilpums 200 cm³ — "
               "aprēķini blīvumu.", "size": 17, "color": GREY}],
        autofit=False)

    lw_ = 3.90
    panel(s, MX, 1.60, lw_, 3.86, [
        {"t": "Dots:", "size": 22, "bold": True, "color": NAVY},
        {"t": "m = 540 g", "size": 19, "space": 7},
        {"t": "V = 200 cm³", "size": 19, "space": 4},
        {"t": "Jāaprēķina:", "size": 22, "bold": True, "color": NAVY,
         "space": 18},
        {"t": "ρ = ?", "size": 19, "space": 7},
        {"t": "Formulas:", "size": 22, "bold": True, "color": NAVY,
         "space": 18},
        {"t": "ρ = m / V", "size": 19, "space": 7},
    ], accent=LINEGREY, fill=LIGHTGREY)

    panel(s, MX + lw_ + 0.30, 1.60, CW - lw_ - 0.30, 3.86, [
        {"t": "Aprēķins:", "size": 22, "bold": True, "color": BLUE},
        {"t": "1)  m = 540 g = 0,54 kg", "size": 19, "space": 10},
        {"t": "2)  V = 200 cm³ = 200 · 10⁻⁶ m³ = 2 · 10⁻⁴ m³", "size": 19,
         "space": 8},
        {"t": "3)  ρ = m / V = 0,54 kg : (2 · 10⁻⁴ m³) = 2700 kg/m³",
         "size": 19, "space": 8},
        {"t": "Atbilde:  ρ = 2700 kg/m³ = 2,7 · 10³ kg/m³", "size": 21,
         "bold": True, "color": RED, "space": 22},
    ], accent=BLUE)

    panel(s, MX, 5.62, CW, 1.16, [
        {"t": "5 SOĻI KATRĀ UZDEVUMĀ", "size": 12, "bold": True,
         "color": GOLD},
        {"t": "1. Dots  →  2. Jāaprēķina  →  3. Formulas  →  4. Aprēķins  "
              "→  5. Atbilde", "size": 20, "bold": True, "space": 7,
         "color": NAVY},
        {"t": "Vispirms pārvērs lielumus SI vienībās, tikai tad ievieto "
              "skaitļus formulā!", "size": 16, "space": 6, "color": GREY},
    ], accent=GOLD, fill=LIGHTGOLD)
    return s


def slide_divider(prs):
    s = blank(prs)
    box(s, 0, 0, SW, SH, fill=NAVY, rounded=False)
    put(s, 1.05, 2.55, 11.0, 1.15,
        [{"t": "UZDEVUMI", "size": 56, "bold": True, "color": WHITE}])
    rule(s, 1.05, 3.90, 3.20, GOLD, 0.03)
    put(s, 1.05, 4.25, 11.0, 1.60, [
        {"t": "Mikropasaule · Makropasaule · Megapasaule", "size": 24,
         "color": RGBColor(0xBD, 0xD7, 0xEE)},
        {"t": "Katrs jaunais solis vispirms parādās liels — pēc tam nostājas "
              "savā vietā risinājumā.", "size": 17, "space": 16,
         "color": RGBColor(0x9D, 0xC3, 0xE6)},
    ])
    return s


# ==================================================================== UZDEVUMI
UZDEVUMI = [
    dict(
        nr=1,
        virsraksts="Ūdens molekulas izmērs · mikropasaule",
        teksts="Ūdens molekulas diametrs ir aptuveni 0,28 nm.\n"
               "a) Izsaki diametru metros standartformā!\n"
               "b) Cik molekulu, cieši novietotu rindā, veidotu 1 mm garu "
               "virkni?",
        dots=["d = 0,28 nm", "L = 1 mm"],
        jaaprekina=["d = ?  (m)", "N = ?"],
        formulas=["N = L / d"],
        aprekins=[
            "1)  d = 0,28 nm = 0,28 · 10⁻⁹ m = 2,8 · 10⁻¹⁰ m",
            "2)  L = 1 mm = 1 · 10⁻³ m",
            "3)  N = L / d = (1 · 10⁻³ m) : (2,8 · 10⁻¹⁰ m) ≈ 3,6 · 10⁶",
        ],
        atbilde="d = 2,8 · 10⁻¹⁰ m ;   N ≈ 3,6 · 10⁶ molekulu",
        piezime="1 mm garā rindā ietilpst ~3,6 miljoni molekulu — "
                "tas ir mikropasaules mērogs.",
    ),
    dict(
        nr=2,
        virsraksts="Detaļas blīvums · makropasaule",
        teksts="Metāla detaļas masa ir 540 g, bet tilpums 200 cm³.\n"
               "Aprēķini vielas blīvumu SI vienībās un nosaki, vai detaļa ir "
               "no alumīnija (ρ(Al) = 2700 kg/m³)!",
        dots=["m = 540 g", "V = 200 cm³", "ρ(Al) = 2700 kg/m³"],
        jaaprekina=["ρ = ?"],
        formulas=["ρ = m / V"],
        aprekins=[
            "1)  m = 540 g = 540 · 10⁻³ kg = 0,54 kg",
            "2)  V = 200 cm³ = 200 · 10⁻⁶ m³ = 2 · 10⁻⁴ m³",
            "3)  ρ = m / V = 0,54 kg : (2 · 10⁻⁴ m³) = 2700 kg/m³",
        ],
        atbilde="ρ = 2700 kg/m³ = 2,7 · 10³ kg/m³ — detaļa ir no alumīnija",
        piezime="Blīvums ir vielas raksturlielums: pēc tā var noteikt, kāda "
                "viela ir ķermenī.",
    ),
    dict(
        nr=3,
        virsraksts="Saules gaisma ceļā uz Zemi · megapasaule",
        teksts="Vidējais attālums no Zemes līdz Saulei ir 1,50 · 10¹¹ m "
               "(1 au), gaismas ātrums vakuumā c = 3,00 · 10⁸ m/s.\n"
               "Cik ilgā laikā Saules gaisma sasniedz Zemi? Izsaki atbildi "
               "arī minūtēs!",
        dots=["s = 1,50 · 10¹¹ m", "c = 3,00 · 10⁸ m/s"],
        jaaprekina=["t = ?  (s)", "t = ?  (min)"],
        formulas=["v = s / t", "t = s / c"],
        aprekins=[
            "1)  t = s / c = (1,50 · 10¹¹ m) : (3,00 · 10⁸ m/s)",
            "2)  t = 0,50 · 10³ s = 5,00 · 10² s = 500 s",
            "3)  t = 500 s : 60 s/min ≈ 8,3 min",
        ],
        atbilde="t = 500 s ≈ 8,3 min  (aptuveni 8 min 20 s)",
        piezime="Mēs redzam Sauli tādu, kāda tā bija pirms ~8 minūtēm.",
    ),
    dict(
        nr=4,
        virsraksts="Zemes vidējais blīvums",
        teksts="Zemes masa ir 5,97 · 10²⁴ kg, bet rādiuss 6,37 · 10⁶ m.\n"
               "Aprēķini Zemes vidējo blīvumu! Lodes tilpumu aprēķina pēc "
               "formulas V = (4/3)·π·R³.",
        dots=["m = 5,97 · 10²⁴ kg", "R = 6,37 · 10⁶ m", "π ≈ 3,14"],
        jaaprekina=["ρ = ?"],
        formulas=["ρ = m / V", "V = (4/3) · π · R³"],
        aprekins=[
            "1)  R³ = (6,37 · 10⁶ m)³ = 2,58 · 10²⁰ m³",
            "2)  V = (4/3) · 3,14 · 2,58 · 10²⁰ m³ ≈ 1,08 · 10²¹ m³",
            "3)  ρ = m / V = (5,97 · 10²⁴ kg) : (1,08 · 10²¹ m³)",
            "4)  ρ ≈ 5,53 · 10³ kg/m³",
        ],
        atbilde="ρ ≈ 5,5 · 10³ kg/m³ = 5500 kg/m³",
        piezime="Vairāk nekā akmens blīvums (~2700 kg/m³), jo Zemes kodols "
                "ir no dzelzs un niķeļa.",
    ),
    dict(
        nr=5,
        virsraksts="Protons un elektrons · mikropasaule",
        teksts="Elektrona masa ir 9,1 · 10⁻³¹ kg, bet protona masa "
               "1,67 · 10⁻²⁷ kg.\n"
               "Cik reižu protona masa ir lielāka par elektrona masu?",
        dots=["mₑ = 9,1 · 10⁻³¹ kg", "mₚ = 1,67 · 10⁻²⁷ kg"],
        jaaprekina=["n = ?"],
        formulas=["n = mₚ / mₑ"],
        aprekins=[
            "1)  n = (1,67 · 10⁻²⁷ kg) : (9,1 · 10⁻³¹ kg)",
            "2)  n = (1,67 : 9,1) · 10⁴ = 0,1835 · 10⁴",
            "3)  n ≈ 1,8 · 10³ ≈ 1835",
        ],
        atbilde="n ≈ 1,8 · 10³ — protons ir ~1800 reižu smagāks",
        piezime="Tāpēc gandrīz visa atoma masa ir kodolā, kaut arī kodols ir "
                "~10 000 reižu mazāks par atomu.",
    ),
    dict(
        nr=6,
        virsraksts="Attālums līdz tuvākajai zvaigznei · megapasaule",
        teksts="Zvaigzne Proksima Kentaura atrodas 4,24 gaismas gadu attālumā "
               "no Zemes.\n"
               "Izsaki šo attālumu metros un parsekos! Izmanto datu bukletu.",
        dots=["s = 4,24 ly"],
        jaaprekina=["s = ?  (m)", "s = ?  (pc)"],
        formulas=["1 ly = 9,46 · 10¹⁵ m", "1 pc = 3,09 · 10¹⁶ m"],
        aprekins=[
            "1)  s = 4,24 · 9,46 · 10¹⁵ m = 40,1 · 10¹⁵ m",
            "2)  s = 4,01 · 10¹⁶ m",
            "3)  s = (4,01 · 10¹⁶ m) : (3,09 · 10¹⁶ m) = 1,30 pc",
        ],
        atbilde="s ≈ 4,01 · 10¹⁶ m ≈ 1,30 pc",
        piezime="Gaisma no šīs zvaigznes līdz mums ceļo vairāk nekā 4 gadus.",
    ),
]

# ------------------------------------------------- uzdevumu slaidu izkārtojums
TOP_Y, TOP_H = 0.70, 1.48                 # aktīvais uzdevums
GAP_Y, GAP_H = 2.26, 3.76                 # ATSTARPE - risinājuma laukums
BOT_Y, BOT_H, BOT_GAP = 6.10, 0.64, 0.06  # nākamie uzdevumi

# Latviešu fizikas uzdevuma standarta izkārtojums:
#   kreisajā šaurajā slejā - Dots: / Jāaprēķina: / Formulas: (ar atdalošām
#   līnijām), vertikāla svītra, labajā pusē - Aprēķins: un Atbilde.
LEFT_X, LEFT_W = MX, 3.30
DIV_X = 4.02
RIGHT_X = 4.24
RIGHT_W = MX + CW - RIGHT_X

CARD_MIN_H = 0.98          # parastais jaunā ieraksta kartītes augstums
CARD_MIN_H_ANS = 1.55      # atbildei atvēlam vairāk vietas
CARD_BOTTOM = GAP_Y + GAP_H - 0.08


L_LABEL, L_ITEM, L_SECGAP = 17.0, 18.0, 16.0   # kreisās slejas tipogrāfija


def _left_lines(sections):
    """Kreisā sleja vienā sarakstā - lai varētu atrast kopīgu mērogu."""
    out = []
    for i, (label, items) in enumerate(sections):
        out.append({"t": label, "size": L_LABEL, "bold": True, "color": NAVY,
                    "space": 0 if i == 0 else L_SECGAP})
        for it in items:
            out.append({"t": it, "size": L_ITEM, "space": 3,
                        "math": MF.has_math(it)})
    return out


def _lines_h(lines, w):
    """Bloka augstums, ņemot vērā, ka daļas ir augstākas par rindu."""
    total = 0.0
    for ln in lines:
        if ln.get("math"):
            total += (math_h(ln["t"], ln["size"], atoms=math_atoms(ln["t"]))
                      + ln.get("space", 0) / 72.0)
        else:
            total += est_h([ln], w)
    return total


def _put_lines(slide, x, y, w, lines, color=DARK, name=None):
    """Uzzīmē sagatavotas rindas; daļas - vertikāli. Atgriež apakšējo malu."""
    for ln in lines:
        sp = ln.get("space", 0) / 72.0
        if ln.get("math"):
            at = math_atoms(ln["t"])
            hh = math_h(ln["t"], ln["size"], atoms=at)
            put_math(slide, x, y + sp, w, hh, ln["t"], ln["size"],
                     color=ln.get("color", color),
                     bold=ln.get("bold", False),
                     align=ln.get("align", PP_ALIGN.LEFT), name=name,
                     atoms=at)
        else:
            hh = est_h([ln], w) - sp
            put(slide, x, y + sp, w, hh, [dict(ln, space=0)], autofit=False,
                name=name)
        y += sp + hh
    return y


def draw_top_task(slide, u, big=False):
    """big=True - uzdevumu tikko parāda: teksts aizņem visu brīvo laukumu."""
    h = (GAP_Y + GAP_H - TOP_Y) if big else TOP_H
    box(slide, MX, TOP_Y, CW, h, fill=LIGHTGOLD, line=GOLD, lw=2.0,
        name="BOX:TASK")
    lines = [{"t": "%d. uzdevums · %s" % (u["nr"], u["virsraksts"]),
              "size": 17 if big else 13, "bold": True, "color": GOLD}]
    for i, part in enumerate(u["teksts"].split("\n")):
        lines.append({"t": part, "size": 34 if big else 20,
                      "space": (16 if i == 0 else 10) if big
                      else (7 if i == 0 else 3)})
    put(slide, MX + 0.30, TOP_Y + 0.14, CW - 0.60, h - 0.28, lines,
        anchor=MSO_ANCHOR.MIDDLE if big else MSO_ANCHOR.TOP,
        name="TXT:TASK")


def draw_bottom_task(slide, u, y):
    n = uid("NEXT")
    box(slide, MX, y, CW, BOT_H, fill=LIGHTGREY, line=LINEGREY, lw=0.75,
        name="BOX:" + n)
    put(slide, MX + 0.24, y + 0.09, CW - 0.48, BOT_H - 0.18,
        [{"t": "%d. uzdevums.  %s" % (u["nr"],
                                      u["teksts"].replace("\n", " ")),
          "size": 13, "color": GREY}], name="TXT:" + n)


def draw_gap(slide, u, step):
    """1 Dots;  2 Jāaprēķina;  3 Formulas;  4..3+k aprēķina soļi;
       4+k Atbilde.  (0. solī rāda tikai uzdevumu - sk. draw_top_task.)"""
    k = len(u["aprekins"])

    # --- kas jau ir "uzrakstīts" un stāv savā vietā -------------------------
    sections = []
    if step >= 2:
        sections.append(("Dots:", u["dots"]))
    if step >= 3:
        sections.append(("Jāaprēķina:", u["jaaprekina"]))
    if step >= 4:
        sections.append(("Formulas:", u["formulas"]))

    n_done = max(0, min(step - 4, k))          # cik aprēķina soļu jau vietā
    shown = list(u["aprekins"][:n_done])

    # --- jaunais ieraksts (liels, centrēts kartītē) ------------------------
    if step == 1:
        card_label, card_items, accent = "DOTS", u["dots"], BLUE
    elif step == 2:
        card_label, card_items, accent = "JĀAPRĒĶINA", u["jaaprekina"], BLUE
    elif step == 3:
        card_label, card_items, accent = "FORMULAS", u["formulas"], BLUE
    elif step <= 3 + k:
        card_label = "APRĒĶINS · %d. SOLIS" % (step - 3)
        card_items, accent = [u["aprekins"][step - 4]], BLUE
    else:
        card_label, card_items, accent = "ATBILDE", [u["atbilde"]], RED

    # --- kreisā sleja: Dots: / Jāaprēķina: / Formulas: ---------------------
    # Mērogu un vietas rēķina pēc PILNA saraksta, lai ieraksts, vienreiz
    # nostājies savā vietā, vairs nekustas un nemaina izmēru.
    inner_w = LEFT_W - 0.10
    full = [("Dots:", u["dots"]), ("Jāaprēķina:", u["jaaprekina"]),
            ("Formulas:", u["formulas"])]
    avail = GAP_H - 0.24 - 0.44
    scale = 1.0
    while scale > 0.34:
        probe = [dict(l, size=l["size"] * scale,
                      space=l.get("space", 0) * scale)
                 for l in _left_lines(full)]
        if _lines_h(probe, inner_w) <= avail:
            break
        scale -= 0.03

    y = GAP_Y + 0.12
    for i, (label, items) in enumerate(full):
        lines = [{"t": label, "size": L_LABEL * scale, "bold": True,
                  "color": NAVY}]
        for it in items:
            lines.append({"t": it, "size": L_ITEM * scale, "space": 3 * scale,
                          "math": MF.has_math(it)})
        h = _lines_h(lines, inner_w)
        if i < len(sections):                       # jau atklāts
            _put_lines(slide, LEFT_X, y, inner_w, lines,
                       name="TXT:SOLVE%d" % i)
        y += h
        if i < 2:
            y += 0.10
            if i + 1 < len(sections):               # svītra tikai starp
                rule(slide, LEFT_X, y, LEFT_W - 0.35)
            y += 0.12

    if sections:
        # vertikālā svītra starp doto un aprēķinu
        box(slide, DIV_X, GAP_Y + 0.06, 0.014, GAP_H - 0.12,
            fill=LINEGREY, line=None, rounded=False)

    apr_y = GAP_Y + 0.10

    # --- aprēķina soļi; ja pietrūkst vietas, noņem vecākos -----------------
    def apr_lines(items, trimmed):
        # Jau atklātie soļi jāzīmē tāpat kā "Formulas:" sleja - ar daļām un
        # saknēm, nevis rindas tekstā (rules_lessons.txt).
        out = [{"t": "Aprēķins:", "size": 17, "bold": True, "color": NAVY}]
        if trimmed:
            out.append({"t": "…", "size": 16, "color": GREY, "space": 4})
        for it in items:
            out.append({"t": it, "size": 17, "space": 6,
                        "math": math_line(it)})
        return out

    # ja soļiem pietrūkst vietas, noņem vecākos; ja neviens vairs neietilpst,
    # aprēķina sarakstu nerāda vispār (nevis vientuļu "...")
    # kamēr kreisā sleja vēl tukša, jaunais ieraksts iznāk pāri visam platumam
    card_x, card_w = (RIGHT_X, RIGHT_W) if sections else (MX, CW)
    apr_w = RIGHT_W - 0.10

    min_card = CARD_MIN_H_ANS if step == 4 + k else CARD_MIN_H
    trimmed = 0
    while True:
        lines = apr_lines(shown, trimmed) if shown else []
        apr_h = _lines_h(lines, apr_w) if lines else 0.0
        card_top = apr_y + (apr_h + 0.22 if lines else 0.0)
        if CARD_BOTTOM - card_top >= min_card or not shown:
            break
        shown.pop(0)
        trimmed += 1

    if lines:
        _put_lines(slide, RIGHT_X, apr_y, apr_w, lines, name="TXT:CALC")

    # --- uzmanības kartīte -------------------------------------------------
    wanted_top = CARD_BOTTOM - max(min_card, CARD_BOTTOM - card_top)
    floor_top = (apr_y + apr_h + 0.16) if lines else (GAP_Y + 0.10)
    card_top = max(wanted_top, floor_top)      # nekad nepārklāj aprēķinu
    card_h = CARD_BOTTOM - card_top
    fill = LIGHTRED if accent is RED else LIGHTBLUE
    box(slide, card_x, card_top, card_w, card_h, fill=fill, line=accent,
        lw=2.0, name="BOX:CARD")

    clines = [{"t": card_label, "size": 13, "bold": True, "color": accent,
               "align": PP_ALIGN.CENTER}]
    for it in card_items:
        clines.append({"t": it, "size": 32, "bold": True, "color": NAVY,
                       "align": PP_ALIGN.CENTER, "space": 10,
                       "math": MF.has_math(it)})
    if step == 4 + k and u.get("piezime"):
        clines.append({"t": u["piezime"], "size": 15, "italic": True,
                       "color": GREY, "align": PP_ALIGN.CENTER, "space": 12})

    cx, cy = card_x + 0.26, card_top + 0.14
    cw, ch = card_w - 0.52, card_h - 0.28
    if any(l.get("math") for l in clines):
        # samazina, līdz visas rindas (arī vertikālās daļas) ietilpst
        sc = 1.0
        while sc > 0.34:
            probe = [dict(l, size=l["size"] * sc,
                          space=l.get("space", 0) * sc) for l in clines]
            if (_lines_h(probe, cw) <= ch
                    and all(not l.get("math")
                            or math_w(l["t"], l["size"], l.get("bold", False))
                            <= cw for l in probe)):
                break
            sc -= 0.04
        probe = [dict(l, size=l["size"] * sc, space=l.get("space", 0) * sc)
                 for l in clines]
        y0 = cy + max(0.0, (ch - _lines_h(probe, cw)) / 2.0)
        _put_lines(slide, cx, y0, cw, probe, name="TXT:CARD")
    else:
        put(slide, cx, cy, cw, ch, clines, anchor=MSO_ANCHOR.MIDDLE,
            name="TXT:CARD")


def task_slides(prs, idx):
    u = UZDEVUMI[idx]
    k = len(u["aprekins"])
    labels = (["Izlasi uzdevumu", "Dots", "Jāaprēķina", "Formulas"]
              + ["Aprēķins · %d. solis" % (i + 1) for i in range(k)]
              + ["Atbilde"])
    for step in range(5 + k):
        s = blank(prs)
        put(s, MX, 0.22, 8.6, 0.30,
            [{"t": "%d. uzdevums no %d" % (u["nr"], len(UZDEVUMI)),
              "size": 14, "bold": True, "color": NAVY}], autofit=False)
        put(s, 9.2, 0.22, CW - 8.65, 0.30,
            [{"t": labels[step], "size": 14, "bold": True, "color": GOLD,
              "align": PP_ALIGN.RIGHT}], autofit=False)
        rule(s, MX, 0.58, CW, GOLD, 0.022)

        draw_top_task(s, u, big=(step == 0))
        if step > 0:
            draw_gap(s, u, step)

        y = BOT_Y
        for nxt in range(idx + 1, min(idx + 3, len(UZDEVUMI))):
            draw_bottom_task(s, UZDEVUMI[nxt], y)
            y += BOT_H + BOT_GAP
        if idx + 1 >= len(UZDEVUMI):
            put(s, MX, BOT_Y + 0.20, CW, 0.34,
                [{"t": "Šis bija pēdējais uzdevums.", "size": 14,
                  "color": GREY, "align": PP_ALIGN.CENTER}], autofit=False)


def slide_kopsavilkums(prs):
    s = blank(prs)
    header(s, "Kopsavilkums un mājasdarbs")

    colw = (CW - 0.38) / 2
    panel(s, MX, 1.20, colw, 3.94, [
        {"t": "ŠODIEN IEMĀCĪJĀMIES", "size": 12, "bold": True,
         "color": BLUE},
        {"t": "•  Matērija pastāv kā VIELA (masa) un kā LAUKS (stiprums).",
         "size": 17, "space": 12},
        {"t": "•  Pasauli nosacīti iedala mikro-, makro- un megapasaulē pēc "
              "objektu izmēra.", "size": 17, "space": 9},
        {"t": "•  Katrai pasaulei ir savas pētīšanas metodes un mērierīces.",
         "size": 17, "space": 9},
        {"t": "•  Lielumus pirms aprēķina izsaka SI vienībās, lietojot "
              "priedēkļus.", "size": 17, "space": 9},
        {"t": "•  Risinājumu pieraksta: Dots — Jāaprēķina — Formulas — "
              "Aprēķins — Atbilde.", "size": 17, "space": 9},
    ], accent=BLUE)

    panel(s, MX + colw + 0.38, 1.20, colw, 3.94, [
        {"t": "MĀJASDARBS", "size": 12, "bold": True, "color": GOLD},
        {"t": "1.  Izsaki SI vienībās:", "size": 17, "space": 12},
        {"t": "      a) 45 nm      b) 2,5 µm      c) 380 g", "size": 17,
         "space": 6},
        {"t": "      d) 1,5 km     e) 0,75 cm³", "size": 17, "space": 4},
        {"t": "2.  Vara gabala masa ir 89 g, tilpums 10 cm³. Aprēķini "
              "blīvumu! Kāda viela tā ir?", "size": 17, "space": 12},
        {"t": "3.  Mēness attālums no Zemes ir 3,85 · 10⁸ m. Cik ilgā laikā "
              "radiosignāls (v = c) nokļūst līdz Mēnesim?", "size": 17,
         "space": 12},
    ], accent=GOLD, fill=LIGHTGOLD)

    panel(s, MX, 5.30, CW, 1.62, [
        {"t": "PAŠVĒRTĒJUMS — atzīmē, cik droši jūties", "size": 12,
         "bold": True, "color": NAVY},
        {"t": "Protu atšķirt vielu no lauka   ·   Protu klasificēt objektu "
              "mikro-/makro-/megapasaulē   ·   Protu pārvērst mērvienības SI "
              "sistēmā   ·   Protu noformēt risinājumu", "size": 17,
         "space": 9},
        {"t": "Nākamā stunda: pasaules pētīšanas metodes un mērierīces — "
              "mikroskops, sensori, teleskops.", "size": 15, "space": 12,
         "color": GREY, "italic": True},
    ], accent=NAVY)
    return s


# ==================================================================== BŪVĒŠANA
def build(out_path):
    prs = new_deck()
    slide_title(prs)
    slide_merkis(prs)
    slide_materija(prs)
    slide_pasaules(prs)
    slide_limeni(prs)
    slide_si(prs)
    slide_priedekli(prs)
    slide_pieraksts(prs)
    slide_divider(prs)
    for i in range(len(UZDEVUMI)):
        task_slides(prs, i)
    slide_kopsavilkums(prs)
    d = os.path.dirname(out_path)
    if d:
        os.makedirs(d, exist_ok=True)
    prs.save(out_path)
    return len(prs.slides._sldIdLst)


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    out = ("C:/aphysics/Dabaszinibas/1. Pasaule ap mums un tās pētīšana/"
           "1.1. Matērija. Viela un lauks. Pasaules līmeņi.pptx")
    print("Izveidots: %s" % out)
    print("Slaidu skaits: %d" % build(out))
