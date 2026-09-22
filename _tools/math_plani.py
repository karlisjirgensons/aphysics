# -*- coding: utf-8 -*-
"""Matemātikas 1.-9. klases tematu un stundu plānu dzinējs.

Paraugs: Math/plan_example.docx (fizikas tematu plāns). Atšķirība - matemātikā
stunda ir katru mācību dienu, laboratorijas darbu nav, un katrs temats
noslēdzas ar vienu pārbaudes darbu (rules_matematika.txt).

Šis modulis atbild tikai par to, KĀ plāns izskatās un kā stundas sakrīt ar
kalendāru (SRP). KO māca, zina klases faili math_1.py ... math_9.py, un tie ir
vienīgā satura vieta - no tiem aug gan plāna dokuments, gan vietnes lapas
(DRY).

Datu modelis:
    Stunda  - viena mācību stunda: jautājums un sasniedzamais rezultāts.
    Bloks   - 3-10 stundu mikrotemats; ar to plānā ir makro atstatums un
              vietnē - navigācija.
    Temats  - programmas temats (mat_p.pdf) ar blokiem un vienu PD.
    Klase   - viena klase: temati + mācību gada noslēguma bloks.

Lietošana:
    import math_plani
    math_plani.klase(3).dokuments("C:/aphysics/Math/math_3.docx")
"""

import datetime as dt
import importlib
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

# Dokumenta pamatelementi ir tie paši, kas fizikas plāniem - viens izskats
# visiem plāniem vienā vietnē (DRY). Fizikas modulis importējot neko nedara.
from fiz_plani import GREY, NAVY, WHITE, cell_text, para, shade

SAKNE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAPE = os.path.join(SAKNE, "Math")

PRIEKSMETS = "Matemātika"
GADS = "2026./2027."
RITMS = "katru mācību dienu - viena stunda dienā, piecas nedēļā"

# ------------------------------------------------------------ mācību kalendārs
SAKUMS = dt.date(2026, 9, 1)
BEIGAS = dt.date(2027, 5, 31)

BRIVLAIKI = [(dt.date(2026, 10, 19), dt.date(2026, 10, 23)),
             (dt.date(2026, 12, 23), dt.date(2027, 1, 4)),
             (dt.date(2027, 3, 15), dt.date(2027, 3, 19))]
SVETKI = [dt.date(2026, 11, 18),      # Latvijas Republikas proklamēšanas diena
          dt.date(2027, 3, 26),       # Lielā Piektdiena
          dt.date(2027, 3, 29),       # Otrās Lieldienas
          dt.date(2027, 5, 4)]        # Neatkarības atjaunošanas diena


def d(x):
    return x.strftime("%d.%m.%Y")


def macibu_dienas(sakums=SAKUMS, beigas=BEIGAS, brivlaiki=BRIVLAIKI,
                  svetki=SVETKI):
    """Visas mācību dienas - darbdienas bez brīvlaikiem un svētku dienām."""
    out, cur = [], sakums
    while cur <= beigas:
        if (cur.weekday() < 5
                and not any(a <= cur <= b for a, b in brivlaiki)
                and cur not in svetki):
            out.append(cur)
        cur += dt.timedelta(days=1)
    return out


def _uzskaite(gabali):
    """«a, b un c» - pēdējo saista ar «un», nevis ar komatu."""
    gabali = list(gabali)
    if len(gabali) < 2:
        return "".join(gabali)
    return "%s un %s" % (", ".join(gabali[:-1]), gabali[-1])


def _posms(a, b):
    """«19.-23.10.2026.», bet pāri mēnešiem «23.12.2026.-04.01.2027.»"""
    if (a.year, a.month) == (b.year, b.month):
        return "%d.-%s." % (a.day, d(b))
    return "%s.-%s." % (d(a), d(b))


def kalendara_apraksts():
    return ("Mācību gada sākums %s., noslēgums %s. Matemātikas stundas notiek "
            "%s. Brīvlaiki: %s Stundu nav %s Datumi ir provizoriski un "
            "aprēķināti pēc šī kalendāra."
            % (d(SAKUMS), d(BEIGAS), RITMS,
               _uzskaite(_posms(a, b) for a, b in BRIVLAIKI),
               _uzskaite("%s." % d(x) for x in SVETKI)))


# ============================================================== satura modelis
MIN_BLOKS, MAX_BLOKS = 3, 10          # makro atstatums (rules_matematika.txt)


class Stunda(object):
    """Viena mācību stunda: jautājums un sasniedzamais rezultāts."""

    __slots__ = ("tema", "sr", "veids", "nr", "datums", "bloks", "temats")

    def __init__(self, tema, sr, veids=None):
        self.tema, self.sr, self.veids = tema, sr, veids
        self.nr = self.datums = self.bloks = self.temats = None

    @property
    def pd(self):
        return self.veids == "PD"


class Bloks(object):
    """Mikrotemats - 3-10 stundas, kas pieder kopā."""

    def __init__(self, nosaukums, stundas):
        self.nosaukums = nosaukums
        self.stundas = [s if isinstance(s, Stunda) else Stunda(*s)
                        for s in stundas]
        for s in self.stundas:
            s.bloks = self

    def __len__(self):
        return len(self.stundas)


class Temats(object):
    """Programmas temats: mikrotematu bloki un noslēguma pārbaudes darbs."""

    def __init__(self, kods, nosaukums, apraksts, bloki, pd, pd_sr,
                 padzilinajumi=None):
        self.kods, self.nosaukums = kods, nosaukums
        self.apraksts, self.padzilinajumi = apraksts, padzilinajumi
        self.bloki = bloki
        self.pd_tema, self.pd_sr = pd, pd_sr
        self.pd = Stunda("%s: %s" % ("PD", pd), pd_sr, "PD")
        self.kods_pd = None           # PD numurs klasē - piešķir Klase
        self.svars = 0                # PD svars procentos - piešķir Klase
        for b in bloki:
            for s in b.stundas:
                s.temats = self
        self.pd.temats = self
        self.pd.bloks = Bloks("Summatīvā vērtēšana", [])

    @property
    def stundas(self):
        """Visas temata stundas pēc kārtas - beigās pārbaudes darbs."""
        return [s for b in self.bloki for s in b.stundas] + [self.pd]

    def __len__(self):
        return len(self.stundas)


def B(nosaukums, stundas):
    """Mikrotemats - satura failiem īsāks vārds."""
    return Bloks(nosaukums, stundas)


def T(kods, nosaukums, apraksts, bloki, pd, pd_sr, padzilinajumi=None):
    """Temats - satura failiem īsāks vārds."""
    return Temats(kods, nosaukums, apraksts, bloki, pd, pd_sr, padzilinajumi)


# ================================================================ klases plāns
class Klase(object):
    """Vienas klases plāns: stundu numuri, datumi un PD svari.

    Saturs par kalendāru nezina - datumus piešķir šeit, tāpēc to pašu plānu
    var izdrukāt citam mācību gadam, nemainot nevienu stundu.
    """

    def __init__(self, nr, modulis):
        self.nr = nr
        self.nosaukums = "%d. klase" % nr
        self.ievads = getattr(modulis, "IEVADS", "")
        self.temati = list(modulis.TEMATI)
        self.noslegums = list(getattr(modulis, "NOSLEGUMS", []))
        self.dienas = macibu_dienas()
        self._numure()
        self._svari()
        self.parbaudi()

    # -- stundu numuri un datumi ------------------------------------------
    @property
    def visas(self):
        return ([s for t in self.temati for s in t.stundas]
                + [s for b in self.noslegums for s in b.stundas])

    def _numure(self):
        for i, s in enumerate(self.visas):
            s.nr = i + 1
            s.datums = self.dienas[i] if i < len(self.dienas) else None

    # -- vērtējumu svari ---------------------------------------------------
    def _svari(self):
        """PD svars ir proporcionāls temata garumam; summa - tieši 100 %."""
        garumi = [len(t) for t in self.temati]
        kopa = sum(garumi)
        svari = [max(5, int(round(100.0 * g / kopa))) for g in garumi]
        svari[-1] += 100 - sum(svari)
        for i, (t, sv) in enumerate(zip(self.temati, svari)):
            t.kods_pd = "PD%d" % (i + 1)
            t.svars = sv
            # Svars stundas virsrakstā neiet: to redz gan PD lapā, gan
            # vērtēšanas kalendārā, un nosaukumā tas tikai pagarina
            # pogu un faila vārdu.
            t.pd.tema = "%s: %s" % (t.kods_pd, t.pd_tema)

    # -- pārbaudes ---------------------------------------------------------
    def parbaudi(self):
        n = len(self.visas)
        assert n <= len(self.dienas), (
            "%d. klasē plānā ir %d stundas, bet mācību gadā tikai %d"
            % (self.nr, n, len(self.dienas)))
        assert n == len(self.dienas), (
            "%d. klasē plānā ir %d stundas, mācību gadā %d - atlikušas %d "
            "neaizpildītas dienas" % (self.nr, n, len(self.dienas),
                                      len(self.dienas) - n))
        svars = sum(t.svars for t in self.temati)
        assert svars == 100, ("%d. klasē PD svaru summa ir %d %%, jābūt 100 %%"
                              % (self.nr, svars))
        for t in self.temati + [None]:
            for b in (t.bloki if t else self.noslegums):
                assert MIN_BLOKS <= len(b) <= MAX_BLOKS, (
                    "%d. klase, «%s»: blokā %d stundas, atļauts %d-%d"
                    % (self.nr, b.nosaukums, len(b), MIN_BLOKS, MAX_BLOKS))

    # -- kopsavilkumi ------------------------------------------------------
    @property
    def stundu_skaits(self):
        return len(self.visas)

    def vertejumi(self):
        """[(kods, tēma, svars, datums, piezīme), ...] - kalendāra tabulai."""
        return [(t.kods_pd, t.pd_tema, t.svars, d(t.pd.datums),
                 "Vērtē tikai stundās apgūto %s. temata saturu." % t.kods[:-1])
                for t in self.temati]

    # -- dokuments ---------------------------------------------------------
    def dokuments(self, path=None):
        path = path or os.path.join(MAPE, "math_%d.docx" % self.nr)
        _dokuments(self, path)
        return path


def _jauns_dokuments(klase):
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    s.left_margin = s.right_margin = Cm(1.2)
    s.top_margin = s.bottom_margin = Cm(1.0)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9)

    para(doc, "%s  |  %s  |  %s  |  5 stundas nedēļā  |  %d mācību stunda"
         % (PRIEKSMETS, klase.nosaukums, GADS, klase.stundu_skaits),
         size=9, bold=True, color=NAVY, after=1)
    para(doc, klase.ievads, size=8.5, color=GREY, after=8)
    return doc


def _kalendars(doc, klase):
    para(doc, "Darba organizācija un vērtēšanas kalendārs", size=13,
         bold=True, color=NAVY, after=3)
    para(doc, kalendara_apraksts(), size=8.5, color=GREY, after=6)
    para(doc, "Katrs temats noslēdzas ar vienu pārbaudes darbu; pārbaudes "
              "darbs aizņem vienu mācību stundu un notiek temata pēdējā "
              "stundā. Pirmais pārbaudes darbs PD1 ir %s."
         % d(klase.temati[0].pd.datums),
         size=8.5, italic=True, color=NAVY, after=6)

    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["Darbs", "Temats / vērtēšanas objekts", "Svars",
                           "Datums", "Piezīme"]):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], h, size=9, bold=True, color=WHITE)
    for kods, tema, svars, datums, piez in klase.vertejumi():
        row = t.add_row().cells
        cell_text(row[0], kods, bold=True, color=NAVY)
        cell_text(row[1], tema)
        cell_text(row[2], "%d %%" % svars, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[3], datums, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[4], piez, size=8, color=GREY)
    for c, w in enumerate([2.0, 9.5, 1.6, 2.8, 11.4]):
        for row in t.rows:
            row.cells[c].width = Cm(w)

    para(doc, "Kopā %d pārbaudes darbi - pa vienam pēc katra temata. Svaru "
              "summa - 100 %%. PD vērtē tikai stundās faktiski apgūto saturu; "
              "neapgūtus izvēles padziļinājumus darbā neiekļauj."
         % len(klase.temati), size=8.5, before=6, after=10)


def _stundu_tabula(doc, stundas):
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    hdr = ["N.p.k.", "Mikrotemats", "Stundas tēma / jautājums",
           "Stundas sasniedzamais rezultāts", "Provizoriskais datums", "St."]
    for i, h in enumerate(hdr):
        shade(t.rows[0].cells[i], "1F3864")
        cell_text(t.rows[0].cells[i], h, size=8.5, bold=True, color=WHITE)
    for s in stundas:
        row = t.add_row().cells
        cell_text(row[0], "%d." % s.nr, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[1], s.bloks.nosaukums, bold=s.pd)
        cell_text(row[2], s.tema, bold=s.pd, color=NAVY if s.pd else None)
        cell_text(row[3], s.sr)
        cell_text(row[4], d(s.datums), align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(row[5], "1", align=WD_ALIGN_PARAGRAPH.CENTER)
        if s.pd:
            for c in row:
                shade(c, "FFF4D6")
    for c, w in enumerate([1.3, 4.0, 6.0, 11.0, 3.0, 1.0]):
        for row in t.rows:
            row.cells[c].width = Cm(w)


def _temata_tabula(doc, temats):
    para(doc, "%s %s (%d stundas)"
         % (temats.kods, temats.nosaukums, len(temats)),
         size=12, bold=True, color=NAVY, before=8, after=2)
    para(doc, temats.apraksts, size=8.5, italic=True, color=GREY, after=2)
    para(doc, "Mikrotemati: %s."
         % "; ".join("%s (%d st.)" % (b.nosaukums, len(b))
                     for b in temats.bloki),
         size=8.5, color=GREY, after=4)
    _stundu_tabula(doc, temats.stundas)
    if temats.padzilinajumi:
        para(doc, "Izvēles padziļinājumi, ja atliek laiks: %s"
             % temats.padzilinajumi, size=8.5, italic=True, color=GREY,
             before=3, after=6)


def _dokuments(klase, path):
    doc = _jauns_dokuments(klase)
    _kalendars(doc, klase)
    for t in klase.temati:
        _temata_tabula(doc, t)
    if klase.noslegums:
        stundas = [s for b in klase.noslegums for s in b.stundas]
        para(doc, "Mācību gada noslēgums (%d stundas)" % len(stundas),
             size=12, bold=True, color=NAVY, before=8, after=2)
        para(doc, "Formatīvas stundas pēc pēdējā pārbaudes darba: apkopo gadā "
                  "apgūto un iezīmē, kas jāatkārto. Jauns summatīvs vērtējums "
                  "nav paredzēts.", size=8.5, italic=True, color=GREY,
             after=4)
        _stundu_tabula(doc, stundas)
    para(doc, "Temati un to secība ņemta no oficiālās programmas parauga "
              "(Math/mat_p.pdf, 4. pielikums); mācību mērķis ir sekmīgs "
              "eksāmens matemātikā 9. klases beigās (Math/mat_ex.pdf).",
         size=8.5, italic=True, color=GREY, before=8)
    doc.save(path)


# ==================================================================== reģistrs
KLASU_SKAITS = 9


def modulis(nr):
    """Klases satura modulis math_<nr>.py."""
    return importlib.import_module("math_%d" % nr)


_KESS = {}


def klase(nr):
    """Vienas klases plāns; vienreiz izrēķināts, tad ņemts no atmiņas."""
    if nr not in _KESS:
        _KESS[nr] = Klase(nr, modulis(nr))
    return _KESS[nr]


def klases():
    """Visas klases, kurām saturs jau uzrakstīts."""
    out = []
    for nr in range(1, KLASU_SKAITS + 1):
        try:
            out.append(klase(nr))
        except ImportError:
            continue
    return out


# ------------------------------------------------------------- faili vietnē
_NEDER = re.compile(r'[<>:"/\\|?*]')


def failam(teksts, garums=70):
    """Nosaukums, ko drīkst likt Windows failā: bez ? : " / \\ | * < >."""
    t = _NEDER.sub("", teksts.replace("«", "").replace("»", ""))
    t = " ".join(t.split()).rstrip(". ")
    return t[:garums].rstrip(". ")


NOSLEGUMA_MAPE = "Mācību gada noslēgums"


def stundas_mape(stunda):
    """Mape, kurā dzīvo stundas lapa; gada noslēgums nepieder nevienam
    tematam, tāpēc tam ir sava mape."""
    return (temata_mape(stunda.temats) if stunda.temats
            else NOSLEGUMA_MAPE)


def temata_mape(temats):
    """Temata mape vietnē: «1.1. Kā izstāsta un parāda - cik, kur, kāds»."""
    return failam("%s %s" % (temats.kods, temats.nosaukums.replace(":", " -")))


def stundas_fails(stunda):
    """Stundas HTML faila nosaukums: «12. Cik kopā, cik palika.html»."""
    return "%d. %s.html" % (stunda.nr, failam(stunda.tema))
